# -*- coding: utf-8 -*-
"""청소년수련시설(여성가족부 청소년수련시설 세부현황) 2020_01(2019.12.31 기준)·2025_01(2024.12.31 기준). 연간 명부 → grade C.
참고: data.go.kr 15030247 20190911판은 내용이 2018.12.31 기준(파일명)이라 쓰지 않고, 여가부 게시판 2019년말 판을 쓴다.
2020년말 판(704791)은 대조용으로 raw에만 둔다."""
import sys, re
from pathlib import Path
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / '_lib'))
import mb
import pandas as pd
import openpyxl
import requests

RAW = HERE / 'raw'
TYP = '청소년수련시설'
DOWN = 'https://www.mogef.go.kr/plc/down.do'
SRC = {
    '2020_01': dict(file='mogef_청소년수련시설_세부현황_20191231.xlsx', bbtSn='704775', atfileSn='705326', ref='2019-12-31', year='2019년말',
                    sheet='현황', ds='여성가족부 2019년도 청소년수련시설 현황(19.12.31 기준) 세부현황'),
    '2025_01': dict(file='mogef_청소년수련시설_세부현황_20241231.xlsx', bbtSn='704849', atfileSn='706150', ref='2024-12-31', year='2024년말',
                    sheet='1.총괄 시설세부현황', ds='여성가족부 2024년 청소년수련시설 현황(24.12.31 기준) 세부현황'),
}


def page(s):
    return f"https://www.mogef.go.kr/mp/pcd/mp_pcd_s001d.do?mid=plc502&bbtSn={s['bbtSn']}"


def download():
    extra = dict(file='mogef_청소년수련시설_세부현황_20201231.xlsx', bbtSn='704791', atfileSn='705486', ref='2020-12-31', ds='2020년말 기준(대조용)')
    for s in list(SRC.values()) + [extra]:
        if (RAW / s['file']).exists():
            continue
        S = requests.Session(); S.headers['User-Agent'] = mb.UA
        S.get(page(s), timeout=60)
        mb.fetch(DOWN, RAW / s['file'], method='POST', session=S, headers={'Referer': page(s)},
                 data=dict(bbid='plc502', bbtSn=s['bbtSn'], atfileSn=s['atfileSn'], atfileSeq='1'), ref_date=s['ref'],
                 note=s['ds'] + ' | 상세페이지 쿠키 후 POST')


def numv(v):
    m = re.search(r'-?\d[\d,]*\.?\d*', str(v or ''))
    try:
        return float(m.group(0).replace(',', '')) if m else None
    except ValueError:
        return None


def read(snap):
    s = SRC[snap]
    ws = openpyxl.load_workbook(RAW / s['file'], read_only=True, data_only=True)[s['sheet']]
    rows = [list(r) for r in ws.iter_rows(values_only=True, max_col=45)]
    h = next(i for i, r in enumerate(rows) if '시설명' in [str(x).strip() if x else '' for x in r])
    H0 = [re.sub(r'\s+', '', str(x or '')) for x in rows[h]]; H1 = [re.sub(r'\s+', '', str(x or '')) for x in rows[h + 1]]
    f0 = lambda k: next(j for j, c in enumerate(H0) if c.startswith(k))
    f1 = lambda k: next((j for j, c in enumerate(H1) if c.startswith(k)), None)
    C = dict(sido=f0('시도'), gu=f0('시군구'), name=f0('시설명'), kind=f0('시설종류'), pub=f0('공공/민간'), founder=f0('설치주체'),
             addr=f0('주소'), cap=f0('수용정원') if any(c.startswith('수용정원') for c in H0) else f0('시설수용정원'),
             site=f1('부지면적'), floor=f1('연면적'), reg=f0('최초등록연월일'))
    rest = next((j for j, c in enumerate(H0) if c.startswith('휴지여부')), None)
    recs = []
    for i in range(h + 2, len(rows)):
        r = rows[i]
        nm = mb.clean(r[C['name']]) if r[C['name']] else ''
        if not nm or '서울' not in str(r[C['sido']] or ''):
            continue
        recs.append(dict(name=nm, gu=mb.clean(r[C['gu']]), facility_subtype=mb.clean(r[C['kind']]), sz_public_private=mb.clean(r[C['pub']]),
                         sz_founder=mb.clean(r[C['founder']]), raw_address=mb.clean(r[C['addr']]), sz_capacity=numv(r[C['cap']]),
                         sz_site_m2=numv(r[C['site']]) if C['site'] is not None else None,
                         sz_area_m2=numv(r[C['floor']]) if C['floor'] is not None else None,
                         sz_registered=str(r[C['reg']] or '')[:10], sz_suspended=mb.clean(r[rest]) if rest is not None else '',
                         source_row_id=f"{s['sheet']}!{i + 1}"))
    d = pd.DataFrame(recs)
    d['address'] = [a if a.startswith('서울') else mb.seoulize(a, g) for a, g in zip(d.raw_address, d.gu)]
    return d


def official(snap):
    ws = openpyxl.load_workbook(RAW / SRC[snap]['file'], read_only=True, data_only=True)['총괄표']
    rows = [list(r) for r in ws.iter_rows(values_only=True, max_col=60)]
    yr = SRC[snap]['year']
    j0 = next((j for j, x in enumerate(rows[1]) if x and yr in str(x)), None)
    if j0 is None:
        return None
    hdr = rows[2]
    js = next(j for j in range(j0, len(hdr)) if hdr[j] and '시' in str(hdr[j]) and '도' in str(hdr[j]))
    names = [re.sub(r'\s+', '', str(hdr[j] or '')) for j in range(js, js + 10)]
    for r in rows[3:]:
        if r[js] and '서울' in str(r[js]) and str(r[js + 1]).strip() == '계':
            return {names[k]: r[js + k] for k in range(2, 9)}
    return None


def build():
    download()
    qa = dict(type=TYP, grade='C', note_grade='연간 공식 명부(여성가족부 청소년수련시설 세부현황, 12.31 기준)', snapshots={})
    fr = {}
    for snap in ['2020_01', '2025_01']:
        d = read(snap)
        d = mb.geocode_df(d, RAW / 'geocoding')
        fr[snap] = d
        qa['snapshots'][snap] = dict(source_file=SRC[snap]['file'], source_rows_seoul=len(d))
    a, b, st = mb.assign_ids(fr['2020_01'], fr['2025_01'], 'YTH')
    qa['facility_id_rule'] = 'mb.assign_ids(0 명칭+주소키+subtype, 1 명칭+주소키, 2 명칭 유일, 3 주소키+subtype 유일, 4 좌표 50m+명칭 앞4자)'
    qa['id_matching'] = st
    for snap, d in [('2020_01', a), ('2025_01', b)]:
        s = SRC[snap]
        d = mb.spatial(d)
        d['category_group'] = '교육보육'; d['facility_type'] = TYP; d['grade'] = 'C'
        d['source_org'] = '여성가족부(현 성평등가족부)'; d['source_dataset'] = s['ds']; d['source_url'] = page(s)
        d['source_file'] = 'raw/' + s['file']; d['source_reference_date'] = s['ref']
        d['reference_month_delta'] = mb.month_delta(s['ref'], snap)
        d['temporal_reason'] = '연말 기준 세부현황'
        d = mb.finalize(d.drop(columns=['raw_address']), snap, HERE, TYP)
        q = qa['snapshots'][snap]
        q['final_rows'] = len(d); q['subtype_counts'] = d.facility_subtype.value_counts().to_dict()
        q['geocoding'] = mb.geo_qa(d); q['duplicates'] = mb.dup_qa(d)
        off = official(snap)
        q['official_check'] = dict(총괄표_서울=off, total_official=off.get('총계') if off else None, parsed=len(d))
    mb.write_json(HERE / f'qa_{TYP}.json', qa)
    return qa


if __name__ == '__main__':
    q = build()
    for s, v in q['snapshots'].items():
        print(s, v['final_rows'], v['subtype_counts'], v['official_check'], v['geocoding']['coord_method'])
    print(q['id_matching'])
