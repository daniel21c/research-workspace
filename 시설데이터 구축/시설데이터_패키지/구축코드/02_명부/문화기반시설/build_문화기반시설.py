# -*- coding: utf-8 -*-
"""문화기반시설(문체부 전국 문화기반시설 총람) 2020_01('20.1.1 기준)·2025_01('25.1.1 기준) 구축. 기준일 배포본 → grade A.
subtype = 총람 시트(국립도서관·공공도서관·박물관·미술관·문예회관·지방문화원·문화의집·생활문화센터). 문학관(2025만)·지역문화재단(부록)은 제외."""
import sys, re, urllib.parse
from pathlib import Path
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / '_lib'))
import mb
import pandas as pd
import openpyxl

RAW = HERE / 'raw'
TYP = '문화기반시설'
M = 'https://www.mcst.go.kr/servlets/eduport/front/upload/UplDownloadFile'
SRC = {
    '2020_01': dict(file='mcst_총람_2020_20200101.xlsx', name='2020 전국 문화기반시설 총람.xlsx', real='DEPTDATA_20210222014111309399.xlsx',
                    ref='2020-01-01', ds='2020 전국 문화기반시설 총람(2020.1.1 기준)', page='https://www.mcst.go.kr/site/s_policy/dept/deptView.jsp?pSeq=1453&pDataCD=0417000000&pType=02'),
    '2025_01': dict(file='mcst_총람_2025_20250101.xlsx', name='2025 전국 문화기반시설 총람.xlsx', real='DEPTDATA_20251217021949027003.xlsx',
                    ref='2025-01-01', ds='2025 전국 문화기반시설 총람(2025.1.1 기준)', page='https://www.mcst.go.kr/site/s_policy/dept/deptView.jsp?pSeq=2078&pDataCD=0417000000&pType=02'),
}
SHEETS = ['국립도서관', '공공도서관', '박물관', '미술관', '생활문화센터', '문예회관', '지방문화원', '문화의집']
EXCLUDED = ['문학관', '(부록)지역문화재단']


def url(s):
    return f"{M}?pFileName={urllib.parse.quote(s['name'])}&pRealName={s['real']}&pPath=0417000000&pFlag="


def download():
    for s in SRC.values():
        mb.fetch(url(s), RAW / s['file'], ref_date=s['ref'], note=s['ds'] + ' | 게시 페이지 ' + s['page'])


def is_data_start(r):
    c0 = r[0]
    if isinstance(c0, (int, float)) or (isinstance(c0, str) and re.fullmatch(r'\d+', c0.strip())):
        return True
    return any(isinstance(x, str) and re.fullmatch(r'\s*(서울|서울특별시|부산|경기|강원)\s*', x) for x in r[1:3])


def read_sheet(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    start = next(i for i in range(4, len(rows)) if is_data_start(rows[i]))
    H = rows[3:start]
    ncol = max(len(r) for r in rows)
    Hf = []
    for k, r in enumerate(H):
        r = list(r) + [None] * (ncol - len(r))
        if k < len(H) - 1:  # 마지막 머리글 행을 빼고 병합 칸을 오른쪽으로 채움
            last = None
            for j in range(ncol):
                if r[j] not in (None, ''):
                    last = r[j]
                elif last is not None:
                    r[j] = last
        Hf.append(r)
    cols = []
    for j in range(ncol):
        parts = []
        for r in Hf:
            v = re.sub(r'\s+', '', str(r[j])) if r[j] not in (None, '') else ''
            if v and v not in parts:
                parts.append(v)
        cols.append('/'.join(parts) or f'col{j}')
    data = []
    for i in range(start, len(rows)):
        r = list(rows[i]) + [None] * (ncol - len(rows[i]))
        data.append([i + 1] + r)
    return ['_row'] + cols, data


def find(cols, *keys, exclude=()):
    for j, c in enumerate(cols):
        if all(k in c for k in keys) and not any(e in c for e in exclude):
            return j
    return None


def numv(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r'-?\d[\d,]*\.?\d*', str(v))
    try:
        return float(m.group(0).replace(',', '')) if m else None
    except ValueError:
        return None


def parse(snap):
    s = SRC[snap]
    wb = openpyxl.load_workbook(RAW / s['file'], read_only=True, data_only=True)
    out = []; counts = {}
    for sh in SHEETS:
        if sh not in wb.sheetnames:
            continue
        cols, data = read_sheet(wb[sh])
        C = lambda *k, exclude=(): find(cols, *k, exclude=exclude)
        name_c = {'국립도서관': lambda: C('도서관명'), '공공도서관': lambda: C('도서관명'),
                  '박물관': lambda: C('박물관', '명', exclude=('주소',)), '미술관': lambda: C('미술관', '명', exclude=('주소',)),
                  '생활문화센터': lambda: C('생활문화센터명'), '문예회관': lambda: C('시설명'), '지방문화원': lambda: C('문화원명'),
                  '문화의집': lambda: next(j for j, c in enumerate(cols) if c.split('/')[0] in ('문화의집',))}[sh]()
        addr_c = C('주소', exclude=('온라인', '홈페이지'))
        sido_c = C('시도') if C('시도') is not None else C('시·도')
        gu_c = C('시군구')
        sz = {}
        if sh in ('국립도서관', '공공도서관'):
            sz = dict(sz_area_m2=C('시설규모', '건물'), sz_site_m2=C('시설규모', '부지'), sz_seats=C('열람석'),
                      sz_books=C('도서자료(권)', exclude=('연간',)), sz_staff=C('현원계'), sz_open_year=C('개관년도'))
            attr = dict(sz_founder=C('설립주체'))
        elif sh in ('박물관', '미술관'):
            sz = dict(sz_area_m2=C('연면적'), sz_site_m2=C('부지면적'), sz_exhibit_m2=C('전시실면적합계'))
            attr = dict(sz_founder=C('구분', '국립'), sz_registered=C('구분', '종') if C('구분', '종') is not None else C('구분', '등록'),
                        sz_open_date=C('개관연월일'))
        elif sh == '생활문화센터':
            sz = dict(sz_area_m2=C('시설현황', '총면적'))
            attr = dict(sz_operation=C('운영방식'), sz_space_type=C('공간유형'), sz_open_date=C('개관연월일') if C('개관연월일') is not None else C('개관연도'))
        elif sh == '문예회관':
            sz = dict(sz_area_m2=C('연면적'))
            attr = dict(sz_founder=C('건립주체'), sz_open_date=C('개관'))
        elif sh == '지방문화원':
            sz = dict(sz_area_m2=C('시설현황', '계'))
            attr = dict(sz_open_date=C('설립일'))
        elif sh == '문화의집':
            sz = dict(sz_area_m2=C('총면적'), sz_visitors=C('연간', '이용자'))
            attr = dict(sz_open_date=C('개관일'))
        seat_cols = [j for j, c in enumerate(cols) if '객석수' in c] if sh == '문예회관' else []
        n_seoul = 0
        for r in data:
            nm = mb.clean(r[name_c]) if r[name_c] is not None else ''
            ad = mb.clean(r[addr_c]) if addr_c is not None and r[addr_c] is not None else ''
            sd = mb.clean(r[sido_c]) if sido_c is not None and r[sido_c] is not None else ''
            if not nm or nm in ('계', '합계', '소계'):
                continue
            if not ('서울' in sd or ad.startswith('서울')):
                continue
            n_seoul += 1
            rec = dict(name=nm, address=ad, facility_subtype=sh, source_row_id=f"{sh}!{r[0]}",
                       gu=mb.clean(r[gu_c]) if gu_c is not None else '')
            for k, j in sz.items():
                rec[k] = numv(r[j]) if j is not None else None
            if seat_cols:
                rec['sz_seats'] = sum(numv(r[j]) or 0 for j in seat_cols)
            for k, j in attr.items():
                rec[k] = mb.clean(r[j]) if j is not None and r[j] is not None else ''
            out.append(rec)
        counts[sh] = n_seoul
    return pd.DataFrame(out), counts, [x for x in wb.sheetnames]


def build():
    download()
    qa = dict(type=TYP, grade='A', note_grade='기준일(1.1) 배포본', snapshots={}, excluded_sheets=EXCLUDED)
    fr = {}
    for snap in ['2020_01', '2025_01']:
        s = SRC[snap]
        d, counts, sheets = parse(snap)
        d['address'] = [a if a.startswith('서울') or not a else mb.seoulize(a, g) for a, g in zip(d.address, d.gu)]
        d = mb.geocode_df(d, RAW / 'geocoding')
        d = mb.spatial(d)
        d['category_group'] = '문화체육녹지'; d['facility_type'] = TYP; d['grade'] = 'A'
        d['source_org'] = '문화체육관광부'; d['source_dataset'] = s['ds']; d['source_url'] = url(s)
        d['source_file'] = 'raw/' + s['file']; d['source_reference_date'] = s['ref']
        d['reference_month_delta'] = mb.month_delta(s['ref'], snap)
        d['temporal_reason'] = '기준일 배포본(1월 1일 기준)'
        fr[snap] = d
        qa['snapshots'][snap] = dict(source_file=s['file'], sheets_in_file=sheets, source_rows_seoul_by_sheet=counts,
                                     source_rows_seoul=sum(counts.values()))
    a, b, st = mb.assign_ids(fr['2020_01'], fr['2025_01'], 'CUL')
    qa['facility_id_rule'] = 'mb.assign_ids(0 명칭+주소키+subtype, 1 명칭+주소키, 2 명칭 유일, 3 주소키+subtype 유일, 4 좌표 50m 이내+명칭 앞4자). 정규화=괄호·법인격·기호 제거, 주소키=도로명+건물번호(없으면 동+번지)'
    qa['id_matching'] = st
    for snap, d in [('2020_01', a), ('2025_01', b)]:
        d = mb.finalize(d.drop(columns=['gu']), snap, HERE, TYP)
        q = qa['snapshots'][snap]
        q['final_rows'] = len(d); q['subtype_counts'] = d.facility_subtype.value_counts().to_dict()
        q['geocoding'] = mb.geo_qa(d); q['duplicates'] = mb.dup_qa(d)
        q['geocoding_by_subtype'] = d.groupby('facility_subtype')['lon'].apply(lambda x: round(x.notna().mean(), 4)).to_dict()
    lib = HERE.parent / '공공도서관' / 'qa_공공도서관.json'
    if lib.exists():
        import json
        L = json.load(open(lib, encoding='utf-8'))
        qa['official_check'] = {snap: dict(chongram_public_library=qa['snapshots'][snap]['subtype_counts'].get('공공도서관'),
                                           libsta_public_library=L['snapshots'][snap]['final_rows']) for snap in qa['snapshots']}
    qa['official_check_note'] = '총람 xlsx에는 시도별 집계표가 없어 시트 행수가 곧 공식 수. 공공도서관은 국가도서관통계 명부 행수와 대조.'
    mb.write_json(HERE / f'qa_{TYP}.json', qa)
    return qa


if __name__ == '__main__':
    q = build()
    for s, v in q['snapshots'].items():
        print(s, v['source_rows_seoul_by_sheet'], v['final_rows'], v['geocoding']['coord_rate'], v['geocoding_by_subtype'], v['duplicates'])
    print(q['id_matching'], q.get('official_check'))
