# -*- coding: utf-8 -*-
"""공공체육시설(문체부 전국 공공체육시설 현황) 2020_01(2019년말 기준)·2025_01(2024년말 기준) 구축. 연간 명부 → grade C.
개별 시설 시트만 명부로 만들고, 간이운동장(마을체육시설)은 자치구 집계표(간이운동장_구별_*.csv)로 따로 낸다.
주소 칸이 없는 시트(2019 생활체육관·게이트볼장·국궁장 등)는 같은 facility_id로 짝지어진 다른 시점 행의 주소를 빌려 좌표화하고
address_source 열에 표시한다."""
import sys, re, urllib.parse
from pathlib import Path
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / '_lib'))
import mb
sys.path.insert(0, str(HERE))
import coord_fill
import pandas as pd
import openpyxl

RAW = HERE / 'raw'
TYP = '공공체육시설'
M = 'https://www.mcst.go.kr/servlets/eduport/front/upload/UplDownloadFile'
SRC = {
    '2020_01': dict(file='mcst_공공체육시설_2019말_총괄.xlsx', name='붙임1전국공공체육시설(2019년말 기준, 총괄).xlsx',
                    real='DEPTDATA_20210210095351642379.xlsx', ref='2019-12-31', ds='전국 공공체육시설 현황(2019년말 기준, 총괄)',
                    village=dict(file='mcst_공공체육시설_2019말_마을체육시설.xlsx', name='붙임1전국공공체육시설(2019년말 기준, 마을체육시설).xlsx',
                                 real='DEPTDATA_20210210095351436996.xlsx', sheet='서울')),
    '2025_01': dict(file='mcst_공공체육시설_2024말_종합본.xlsx', name='2025 전국 공공체육시설 현황(종합본).xlsx',
                    real='DEPTDATA_20260317100013749945.xlsx', ref='2024-12-31', ds='2025 전국 공공체육시설 현황(2024년말 기준, 종합본)',
                    village=None),
}
SKIP = ('표지', '일반개요', '공공체육시설의분류기준', '연도별현황', '설치주체별현황', '시도별현황', 'Sheet1', '8.총괄표')
SIDO = ['서울', '부산', '대구', '인천', '광주', '대전', '울산', '세종', '경기', '강원', '충북', '충남', '전북', '전남', '경북', '경남', '제주', '전국']


def url(name, real):
    return f"{M}?pFileName={urllib.parse.quote(name)}&pRealName={real}&pPath=0417000000&pFlag="


def download():
    coord_fill.download()
    for s in SRC.values():
        mb.fetch(url(s['name'], s['real']), RAW / s['file'], ref_date=s['ref'], note=s['ds'])
        if s['village']:
            v = s['village']
            mb.fetch(url(v['name'], v['real']), RAW / v['file'], ref_date=s['ref'], note='마을체육시설(간이운동장) ' + s['ds'])


def sido_of(v):
    v = re.sub(r'\s+', '', str(v or ''))
    for s in SIDO:
        if v.startswith(s):
            return s
    return None


def subtype_of(sh):
    s = re.sub(r'^\s*[\d.]+\s*', '', sh).strip()
    s = s.replace('기타 체육시설', '기타체육시설')
    return s


def numv(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r'-?\d[\d,]*\.?\d*', str(v))
    try:
        return float(m.group(0).replace(',', '')) if m else None
    except ValueError:
        return None


def read_sheet(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True, max_col=80)]
    hr = next(i for i, r in enumerate(rows[:6]) if any(re.sub(r'\s+', '', str(x or '')) == '시설명' for x in r))
    ncol = max(len(r) for r in rows[:hr + 4])
    # 시도 열
    H0 = [re.sub(r'\s+', '', str(x or '')) for x in rows[hr]]
    sido_c = H0.index('시도') if '시도' in H0 else None
    end = hr + 1
    while end < min(len(rows), hr + 5):
        r = rows[end]
        if (sido_c is not None and sido_of(r[sido_c]) ) or any(re.fullmatch(r'\s*(계|소\s*계)\s*', str(x or '')) for x in r[:4]):
            break
        end += 1
    Hf = []
    for k in range(hr, end):
        r = list(rows[k]) + [None] * (ncol - len(rows[k]))
        if k < end - 1:
            last = None
            for j in range(ncol):
                if r[j] not in (None, ''):
                    last = r[j]
                elif last is not None:
                    r[j] = last
        Hf.append(r)
    cols = []
    for j in range(ncol):
        parts = []
        for r in Hf:
            v = re.sub(r'\s+', '', str(r[j])) if r[j] not in (None, '') else ''
            if v and v not in parts:
                parts.append(v)
        cols.append('/'.join(parts))
    return cols, rows, end, sido_c


def find(cols, key, exclude=()):
    for j, c in enumerate(cols):
        if key in c and not any(e in c for e in exclude):
            return j
    return None


def parse(snap):
    s = SRC[snap]
    wb = openpyxl.load_workbook(RAW / s['file'], read_only=True, data_only=True)
    recs = []; per_sheet = {}
    for sh in wb.sheetnames:
        if sh in SKIP or re.fullmatch(r'\d\d\S+', sh):  # 01서울 등 간이운동장 시도표
            continue
        cols, rows, start, sido_c = read_sheet(wb[sh])
        name_c = find(cols, '시설명'); addr_c = find(cols, '주소', exclude=('홈페이지',))
        gu_c = find(cols, '시군구') if find(cols, '시군구') is not None else find(cols, '시ㆍ군ㆍ구')
        C = dict(sz_area_m2=find(cols, '부지면적'), sz_building_m2=find(cols, '건축면적'), sz_floor_m2=find(cols, '연면적'),
                 sz_seats=find(cols, '좌석수'), sz_capacity=find(cols, '수용인원'), sz_built_year=find(cols, '준공'))
        A = dict(sz_owner=find(cols, '소유기관'), sz_manager=find(cols, '관리주체') if find(cols, '관리주체') is not None else find(cols, '관리형태'))
        cur = None; n = 0
        for i in range(start, len(rows)):
            r = list(rows[i]) + [None] * (len(cols) - len(rows[i]))
            if sido_c is not None and sido_of(r[sido_c]):
                cur = sido_of(r[sido_c])
            nm = mb.clean(r[name_c]) if name_c is not None and r[name_c] is not None else ''
            gu = mb.clean(r[gu_c]) if gu_c is not None and r[gu_c] is not None else ''
            if cur != '서울' or not nm or re.fullmatch(r'(계|소\s*계|소개)', gu.replace(' ', '')) or re.fullmatch(r'[\d,]+(개소)?', nm):
                continue
            gu = re.sub(r'\s+', '', gu)
            ad = mb.clean(r[addr_c]) if addr_c is not None and r[addr_c] is not None else ''
            rec = dict(name=nm, gu=gu, raw_address=ad, facility_subtype=subtype_of(sh), source_row_id=f'{sh}!{i + 1}',
                       sheet_has_address=addr_c is not None)
            for k, j in C.items():
                rec[k] = numv(r[j]) if j is not None else None
            for k, j in A.items():
                rec[k] = mb.clean(r[j]) if j is not None and r[j] is not None else ''
            recs.append(rec); n += 1
        per_sheet[subtype_of(sh)] = n
    d = pd.DataFrame(recs)
    # 주소 조립: 구 이름 없으면 시군구 열을 붙임
    def compose(ad, gu):
        if not ad:
            return ''
        if ad.startswith('서울') or mb.OTHER_SIDO.match(ad):
            return ad
        if not any(g in ad[:6] for g in mb.GU) and gu in mb.GU:
            return f'서울특별시 {gu} {ad}'
        return '서울특별시 ' + ad
    d['address'] = [compose(a, g) for a, g in zip(d.raw_address, d.gu)]
    return d, per_sheet


def official(snap):
    """시도별현황 시트의 서울 개소·면적."""
    s = SRC[snap]
    ws = openpyxl.load_workbook(RAW / s['file'], read_only=True, data_only=True)['시도별현황']
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    h = [re.sub(r'\s+', '', str(x or '')) for x in rows[0]]
    j = h.index('서울')
    out = {}
    for r in rows[2:]:
        if r[0] and '시설항목' in re.sub(r'\s+', '', str(r[0])):
            break  # 두 번째 블록(경기·강원…)부터는 다른 시도
        if r[0]:
            k = re.sub(r'\s+', '', re.sub(r'^\s*[\d.]+\s*', '', str(r[0]))).lstrip('-')
            out[k] = dict(count=r[j], area=r[j + 1])
    return out


OFF_MAP = {'전천후게이트볼장': ['게이트볼장'], '체육관': ['구기체육관', '투기체육관', '생활체육관'],
           '기타시설': ['기타체육시설', '기타체육시설(풋살장)', '기타체육시설(그외)'], '풋살장': ['기타체육시설(풋살장)'],
           '그외': ['기타체육시설(그외)'], '설상경기장': ['스키점프경기장', '바이애슬론경기장', '크로스컨트리경기장', '봅슬레이,루지,스켈레톤경기장']}


def village(snap):
    s = SRC[snap]
    if s['village']:
        f, sh = RAW / s['village']['file'], s['village']['sheet']
    else:
        f, sh = RAW / s['file'], '01서울'
    ws = openpyxl.load_workbook(f, read_only=True, data_only=True)[sh]
    rows = [list(r) for r in ws.iter_rows(values_only=True, max_col=26)]
    names = ['gu', '시설수_계', '체육공원', '둔치', '마을공터', '아파트단지', '약수터', '등산로', '도시공원', '기타', '시설조성면적_m2',
             '간이운동시설_계', '축구장', '배구장', '농구장', '씨름장', '테니스장', '게이트볼장', '운동광장', '배드민턴장', '체조장', '로울러장',
             '수영장', '기타종목', '체력단련시설_점', '부대편익시설_점']
    out = []
    for r in rows[4:]:
        g = re.sub(r'\s+', '', str(r[0] or ''))
        if g in mb.GU or g in ('계', 'p'):
            rec = dict(zip(names, [g if g != 'p' else '계'] + [numv(x) for x in r[1:26]]))
            out.append(rec)
    df = pd.DataFrame(out)
    # 합계식이 캐시되지 않은 경우 대비: 계열 재계산 검산
    sub = df[df.gu.isin(mb.GU)]
    chk = dict(rows_gu=len(sub), sum_시설수=float(sub['시설수_계'].fillna(0).sum()),
               sum_유형합=float(sub[['체육공원', '둔치', '마을공터', '아파트단지', '약수터', '등산로', '도시공원', '기타']].fillna(0).sum().sum()),
               계_row=df[df.gu == '계']['시설수_계'].tolist())
    return df, chk


def borrow(d_from, d_to):
    """주소 없는 행: 같은 facility_id의 다른 시점 좌표를 빌린다."""
    m = d_from[d_from.lon.notna()].drop_duplicates('facility_id').set_index('facility_id')
    need = d_to.lon.isna() & (d_to.address == '') & d_to.facility_id.isin(m.index)
    for i in d_to.index[need]:
        f = d_to.at[i, 'facility_id']
        d_to.at[i, 'lon'] = m.at[f, 'lon']; d_to.at[i, 'lat'] = m.at[f, 'lat']
        d_to.at[i, 'coord_method'] = m.at[f, 'coord_method']
        d_to.at[i, 'geocode_detail'] = 'borrowed:' + str(m.at[f, 'geocode_detail'])
        d_to.at[i, 'address_source'] = 'borrowed_other_snapshot_same_facility_id:' + m.at[f, 'address']
    return int(need.sum())


def build():
    download()
    REJ = {}; STG = {}
    qa = dict(type=TYP, grade='C', note_grade='연간 공식 명부(연말 기준)', snapshots={})
    fr = {}
    for snap in ['2020_01', '2025_01']:
        s = SRC[snap]
        d, per = parse(snap)
        d = mb.geocode_df(d, RAW / 'geocoding')
        d['address_source'] = ['source' if a else 'none' for a in d.address]
        d, rej, stg = coord_fill.fill(d, snap)   # 1단계: 주소 표기 정리 후 재조회
        fr[snap] = d; REJ[snap] = rej; STG[snap] = stg
        qa['snapshots'][snap] = dict(source_file=s['file'], source_rows_seoul_by_sheet=per, source_rows_seoul=len(d),
                                     rows_without_address=int((d.address == '').sum()))
    # facility_id: subtype+명칭 기준(주소 없는 행이 많아 명칭 가중)
    a, b, st = mb.assign_ids(fr['2020_01'], fr['2025_01'], 'SPO')
    qa['facility_id_rule'] = 'mb.assign_ids(0 명칭+주소키+subtype, 1 명칭+주소키, 2 명칭 유일, 3 주소키+subtype 유일, 4 좌표 50m 이내+명칭 앞4자). 정규화=괄호·법인격·기호 제거, 주소키=도로명+건물번호(없으면 동+번지) 같은 명칭이 여러 종목 시트에 있으면 명칭+주소+종목 규칙이 먼저 적용됨'
    qa['id_matching'] = st
    for x, y, k in [(b, a, '2020_01'), (a, b, '2025_01')]:
        qa['snapshots'][k]['coords_borrowed_from_other_snapshot'] = borrow(x, y)
        y.loc[y.geocode_detail.astype(str).str.startswith('borrowed:'), 'coord_stage'] = 'borrow_other_snapshot_same_id'
    for snap, d in [('2020_01', a), ('2025_01', b)]:
        coord_fill.fill_rest(d, snap, REJ[snap], STG[snap])   # 2단계(공식 명부 차용), 3단계(Kakao 장소 검색)
    for snap, d in [('2020_01', a), ('2025_01', b)]:
        s = SRC[snap]
        d = mb.spatial(d)
        d['category_group'] = '문화체육녹지'; d['facility_type'] = TYP; d['grade'] = 'C'
        d['source_org'] = '문화체육관광부'; d['source_dataset'] = s['ds']; d['source_url'] = url(s['name'], s['real'])
        d['source_file'] = 'raw/' + s['file']; d['source_reference_date'] = s['ref']
        d['reference_month_delta'] = mb.month_delta(s['ref'], snap)
        d['temporal_reason'] = '연말 기준 연간 현황'
        d = mb.finalize(d.drop(columns=['raw_address']), snap, HERE, TYP)
        q = qa['snapshots'][snap]
        q['final_rows'] = len(d); q['subtype_counts'] = d.facility_subtype.value_counts().to_dict()
        q['geocoding'] = mb.geo_qa(d); q['duplicates'] = mb.dup_qa(d)
        q['geocoding_by_subtype'] = d.groupby('facility_subtype')['lon'].apply(lambda x: round(x.notna().mean(), 4)).to_dict()
        q['coord_fill'] = dict(stage_counts=STG[snap], by_stage=d.coord_stage.replace('', 'unresolved').value_counts().to_dict(),
                               by_method=d.coord_method.value_counts().to_dict(), final_coord_rate=round(d.lon.notna().mean(), 4),
                               reject_reasons=REJ[snap],
                               note='1단계 정확 일치 재조회, 2단계 OA-1115 주소/OA-21779 명칭+자치구+종목(좌표 열 없어 그 주소를 정확 지오코딩), '
                                    '3단계 Kakao 키워드(자치구 일치+종목어, 유일 또는 유사도 확실 1위). 2·3단계 좌표는 명칭 대응에 기댄 것이라 오차 가능')
        off = official(snap)
        cmp = {}
        for k, v in off.items():
            if k == '합계' or '간이운동장' in k:
                continue
            types = OFF_MAP.get(k, [k])
            mine = sum(q['subtype_counts'].get(t, 0) for t in types)
            oc = v['count'] if isinstance(v['count'], (int, float)) else 0
            cmp[k] = dict(official_seoul_count=v['count'], parsed_rows=mine, diff=mine - oc)
        q['official_check'] = cmp
        tot = off.get('합계', {}).get('count'); vil = next((v['count'] for k, v in off.items() if '간이운동장' in k), None)
        q['official_total'] = dict(official_total_incl_village=tot, official_village=vil,
                                   official_individual=(tot - vil) if isinstance(tot, (int, float)) and isinstance(vil, (int, float)) else None,
                                   parsed_individual=len(d),
                                   note='체육관 소계는 구기+투기+생활, 기타시설은 풋살장+그외(2024)로 합산 비교')
        vdf, chk = village(snap)
        vdf.insert(0, 'year_snapshot', snap)
        vdf.to_csv(HERE / f'간이운동장_구별_{snap}.csv', index=False, encoding='utf-8-sig')
        q['village_table'] = dict(file=f'간이운동장_구별_{snap}.csv', **chk,
                                  source=(s['village']['file'] + '!서울') if s['village'] else s['file'] + '!01서울')
    mb.write_json(HERE / f'qa_{TYP}.json', qa)
    return qa


if __name__ == '__main__':
    import json
    q = build()
    for s, v in q['snapshots'].items():
        print(s, v['final_rows'], v['rows_without_address'], v.get('coords_borrowed_from_other_snapshot'), v['geocoding']['coord_rate'])
        print('  ', json.dumps(v['official_check'], ensure_ascii=False)[:1500])
        print('  ', v['village_table'])
    print(q['id_matching'])
