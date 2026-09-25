# -*- coding: utf-8 -*-
"""공공도서관(국가도서관통계시스템 개별 도서관 통계) 2020_01·2025_01 구축.
원천: libsta 2019년 실적(→2020_01), 2024년 실적(→2025_01). 연간 명부 → grade C.
실행: python build_공공도서관.py  (이 파일 위치 기준 상대경로)"""
import sys, re
from pathlib import Path
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / '_lib'))
import mb
import pandas as pd
import openpyxl

RAW = HERE / 'raw'
TYP = '공공도서관'
SRC = {
    '2020_01': dict(file='libsta_공공도서관_2019실적.xlsx', url='https://www.libsta.go.kr/libsta/filemanager/FILE_ID_000000000915/4',
                    ref='2019-12-31', ds='국가도서관통계 공공도서관 2019년 실적(개별 도서관)'),
    '2025_01': dict(file='libsta_공공도서관_2024실적.xlsx', url='https://www.libsta.go.kr/libsta/filemanager/FILE_ID_000000001058/8',
                    ref='2024-12-31', ds='국가도서관통계 공공도서관 2024년 실적(2025년 조사, 개별 도서관)'),
}


def download():
    for s in SRC.values():
        mb.fetch(s['url'], RAW / s['file'], ref_date=s['ref'], note=s['ds'])


def num(v):
    if v is None:
        return None
    try:
        return float(str(v).replace(',', '').strip())
    except ValueError:
        return None


def read_2019():
    ws = openpyxl.load_workbook(RAW / SRC['2020_01']['file'], read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    h = [str(x) for x in rows[0]]
    df = pd.DataFrame(rows[1:], columns=h)
    df['_row'] = range(2, len(df) + 2)
    g = lambda c: df[c].map(num)
    out = pd.DataFrame(dict(
        lib_code=df['도서관코드'].astype(str), name=df['도서관명'], subtype=df['구분'], sido=df['지역'], gu=df['지역_시군구'],
        addr=df['주소'], addr_detail=df['상세주소'], row=df['_row'],
        sz_books=sum(g(c) .fillna(0) for c in h if c.startswith('도서_국내_') or c.startswith('도서_국외_')),
        sz_area_m2=g('시설_연면적'), sz_site_m2=g('시설_도서관부지면적'), sz_seats=g('좌석_총좌석수'),
        sz_librarians=g('직원_정규직_사서남').fillna(0) + g('직원_정규직_사서여').fillna(0) + g('직원_비정규직_사서남').fillna(0) + g('직원_비정규직_사서여').fillna(0),
        sz_staff=sum(g(c).fillna(0) for c in h if re.match(r'직원_(정규직_.*(남|여)|비정규직_.*(남|여))$', c)),
        sz_visitors=g('도서관_방문자수'), sz_open_year=g('개관년도'), sz_founder=df['설립주체'], sz_founder_org=df['설립주체_기관명'],
        sz_operation=df['운영방식'], branch=df['분관내용']))
    return out


def read_2024():
    ws = openpyxl.load_workbook(RAW / SRC['2025_01']['file'], read_only=True, data_only=True)['원자료_분석용']
    rows = list(ws.iter_rows(values_only=True))
    cols = ['/' + str(b).replace(chr(10), '') for b in rows[0]]  # '원자료_분석용' 시트: 1행 머리글, 2행부터 자료
    df = pd.DataFrame(rows[1:], columns=cols)
    df['_row'] = range(2, len(df) + 2)
    C = lambda k: next(c for c in cols if c == '/' + k)
    g = lambda k: df[C(k)].map(num)
    out = pd.DataFrame(dict(
        lib_code=df[C('도서관코드')].astype(str), name=df[C('도서관명')], subtype=df[C('구분')], sido=df[C('지역')], gu=df[C('시군구')],
        addr=df[C('주소')], addr_detail=df[C('상세주소')], row=df['_row'],
        sz_books=g('국내서_합계').fillna(0) + g('국외서_합계').fillna(0),
        sz_area_m2=g('면적_도서관 건물 연면적'), sz_site_m2=g('면적_도서관 부지 면적'), sz_seats=g('좌석수_총 좌석수'),
        sz_librarians=g('정규직_사서직_남').fillna(0) + g('정규직_사서직_여').fillna(0) + g('비정규직_사서자격증소지자_남').fillna(0) + g('비정규직_사서자격증소지자_여').fillna(0),
        sz_staff=g('인원수_합계'), sz_visitors=g('이용자수_도서관방문자수'), sz_open_year=g('개관년도'),
        sz_founder=df[C('설립주체')], sz_founder_org=df[C('설립기관명')], sz_operation=df[C('운영방식')], branch=df[C('분관내용')]))
    return out[out.name.notna()]


def chongram_seoul(snap):
    """문화기반시설 총람 공공도서관 시트의 서울 행(교차대조용)."""
    f = HERE.parent / '문화기반시설' / 'raw' / ('mcst_총람_2020_20200101.xlsx' if snap == '2020_01' else 'mcst_총람_2025_20250101.xlsx')
    ws = openpyxl.load_workbook(f, read_only=True, data_only=True)['공공도서관']
    rows = [r for r in ws.iter_rows(min_row=7, values_only=True) if r[1] and '서울' in str(r[1]) and r[4]]
    return pd.DataFrame([dict(name=mb.clean(r[4]), address=mb.clean(r[5])) for r in rows])


def build():
    download()
    frames = {}
    qa = dict(type=TYP, grade='C', note_grade='연간 공식 명부(국가도서관통계 개별 도서관 실적)', snapshots={})
    for snap, reader in [('2020_01', read_2019), ('2025_01', read_2024)]:
        s = SRC[snap]
        d = reader()
        n_all = len(d)
        d = d[d.sido.astype(str).str.contains('서울')].copy()
        n_seoul = len(d)
        d['name'] = d['name'].map(mb.clean)
        # 주소 칸에 건물번호가 없으면(2019판은 번호가 상세주소로 넘어간 행이 많음) 상세주소를 이어 붙인다
        d['address'] = [mb.clean(a) if mb.parse_addr(a) else mb.clean(f"{mb.clean(a)} {mb.clean(x)}")
                        for a, x in zip(d['addr'], d['addr_detail'])]
        d['address'] = [a if a.startswith('서울') else mb.seoulize(a, g) for a, g in zip(d.address, d.gu.map(mb.clean))]
        d = mb.geocode_df(d, RAW / 'geocoding')
        d = mb.spatial(d)
        d['category_group'] = '문화체육녹지'
        d['facility_type'] = TYP
        d['facility_subtype'] = d['subtype'].map(mb.clean)
        d['grade'] = 'C'
        d['source_org'] = '문화체육관광부 국립중앙도서관(국가도서관통계시스템)'
        d['source_dataset'] = s['ds']
        d['source_url'] = s['url']
        d['source_file'] = 'raw/' + s['file']
        d['source_row_id'] = d['row'].astype(str) + ':' + d['lib_code']
        d['source_reference_date'] = s['ref']
        d['reference_month_delta'] = mb.month_delta(s['ref'], snap)
        d['temporal_reason'] = '연간 실적 조사(해당 연도 말 운영 도서관)'
        frames[snap] = d
        qa['snapshots'][snap] = dict(source_file=s['file'], source_rows_national=n_all, source_rows_seoul=n_seoul)
    # facility_id: 도서관코드 우선, 없으면 명칭+주소 규칙
    a, b = frames['2020_01'], frames['2025_01']
    both = set(a.lib_code) & set(b.lib_code)
    a['facility_id'] = 'LIB' + a['lib_code']
    b['facility_id'] = 'LIB' + b['lib_code']
    # 코드가 바뀐 경우(명칭+주소 동일) 2020 id로 통일
    a_rest = a[~a.lib_code.isin(both)]; b_rest = b[~b.lib_code.isin(both)]
    _, b2, st = mb.assign_ids(a_rest.assign(facility_subtype=''), b_rest.assign(facility_subtype=''), 'TMP')
    code_map = {}
    for i, (rid, rule) in enumerate(zip(b2['facility_id'], b2['id_match_rule'])):
        if rule != 'new' and rule in ('name+addr', 'name'):
            j = int(rid[3:]) - 1
            code_map[b_rest.iloc[i]['facility_id']] = a_rest.iloc[j]['facility_id']
    b['facility_id'] = b['facility_id'].map(lambda x: code_map.get(x, x))
    id_rule = '1) 국가도서관통계 도서관코드 동일 → 같은 id(LIB+코드). 2) 코드가 다르면 정규화 명칭+주소키 또는 명칭 유일 일치 시 2020 id 사용.'
    qa['facility_id_rule'] = id_rule
    qa['id_matching'] = dict(same_code=len(both), matched_by_name_after_code_change=len(code_map),
                             only_2020=int((~a.facility_id.isin(b.facility_id)).sum()), only_2025=int((~b.facility_id.isin(a.facility_id)).sum()))
    for snap, d in frames.items():
        d = mb.finalize(d.drop(columns=['addr', 'sido', 'subtype', 'row']), snap, HERE, TYP)
        q = qa['snapshots'][snap]
        q['final_rows'] = len(d)
        q['geocoding'] = mb.geo_qa(d)
        q['duplicates'] = mb.dup_qa(d)
        q['subtype_counts'] = d.facility_subtype.value_counts().to_dict()
        # 총람 공공도서관과 교차대조
        c = chongram_seoul(snap)
        nn = set(d.name.map(mb.norm_name)); cn = set(c.name.map(mb.norm_name))
        ak = set(d.address.map(mb.addr_key)); ca = set(c.address.map(mb.addr_key))
        matched = sum(1 for n, ad in zip(c.name.map(mb.norm_name), c.address.map(mb.addr_key)) if n in nn or ad in ak)
        q['official_check'] = dict(
            chongram_public_library_seoul=len(c), libsta_seoul=len(d), difference=len(d) - len(c),
            chongram_rows_matched_by_name_or_address=matched,
            chongram_rows_unmatched=[n for n, ad in zip(c.name, c.address) if mb.norm_name(n) not in nn and mb.addr_key(ad) not in ak][:40],
            note='총람(1.1 기준)과 국가도서관통계(연말 실적)는 기준일·범위(사립·작은 공공 포함 여부)가 달라 수가 다를 수 있음')
    mb.write_json(HERE / f'qa_{TYP}.json', qa)
    return qa


if __name__ == '__main__':
    import json
    q = build()
    print(json.dumps({k: {kk: vv for kk, vv in v.items() if kk in ('source_rows_seoul', 'final_rows', 'official_check')} for k, v in q['snapshots'].items()}, ensure_ascii=False, indent=1)[:3000])
    print(q['id_matching'])
