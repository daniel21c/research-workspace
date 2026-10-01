# -*- coding: utf-8 -*-
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
OUT = Path(r"D:\Research\00_박사논문_연구체계\01_생활권_필요성\results")
RP = Path(r"D:\Research\00_박사논문_연구체계\01_생활권_필요성\보고_20261002")
md = (RP / "교수님_보고_20261002.md").read_text(encoding="utf-8").splitlines()
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Malgun Gothic"; st.font.size = Pt(10.5); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
for s in doc.sections: s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)
def add_runs(par, text):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"): par.add_run(part[2:-2]).bold = True
        else: par.add_run(part.replace("`", ""))
i = 0; table_rows = []
def flush_table():
    global table_rows
    if not table_rows: return
    rows = [r for r in table_rows if not re.match(r"^\|?\s*-{3,}", r)]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    t = doc.add_table(rows=len(cells), cols=len(cells[0])); t.style = "Table Grid"
    for ri, row in enumerate(cells):
        for ci, c in enumerate(row):
            p = t.cell(ri, ci).paragraphs[0]; add_runs(p, c)
            for run in p.runs: run.font.size = Pt(9); run.bold = run.bold or ri == 0
    table_rows = []
for line in md:
    if line.startswith("|"): table_rows.append(line); continue
    flush_table()
    if line.startswith("# "): doc.add_heading(line[2:].strip(), level=1)
    elif line.startswith("## "): doc.add_heading(line[3:].strip(), level=2)
    elif re.match(r"^\d+\. ", line): add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
    elif line.startswith("- "): add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
    elif line.strip(): add_runs(doc.add_paragraph(), line.strip())
flush_table()
doc.add_heading("그림", level=2)
caps = {"그림_필요성A_생활권지도_2025.png": "그림 A 같은 추가량을 놓기 전·격자 총량 최대·생활권 최저선 5%의 생활권별 6분야 완결 비율(2025)", "그림_필요성B_규칙별빈생활권_2025.png": "그림 B 규칙별로 남는 빈 생활권 수(2025): 시설 하나씩(왼), 6분야 묶음(오)", "그림_묶음1_완결곡선_결손수.png": "그림 C 묶음 정의별 완결률 곡선과 결손 분야 수"}
for f, c in caps.items():
    doc.add_picture(str(OUT / f), width=Cm(17)); p = doc.add_paragraph(c); p.runs[0].font.size = Pt(9)
out = RP / "교수님_보고_20261002.docx"; doc.save(out); print(out, out.stat().st_size)
