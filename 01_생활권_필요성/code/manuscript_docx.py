# -*- coding: utf-8 -*-
"""Build manuscript/manuscript_draft_en.docx from manuscript_draft_en.md, tables/*.md and figures/Fig1–6.
Simple converter: headings (#, ##, ###, ####), paragraphs, bullets, **bold**, *italic*, $math$ kept as text; tables inserted after References as in Elsevier submissions."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
R = Path(__file__).parent.parent / "manuscript"
lines = (R / "manuscript_draft_en.md").read_text(encoding="utf-8").splitlines()
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
for s in doc.sections: s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.5)
def runs(par, text):
    text = text.replace("$", "")
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if part.startswith("**") and part.endswith("**"): par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2: par.add_run(part[1:-1]).italic = True
        elif part.startswith("`"): par.add_run(part[1:-1])
        else: par.add_run(part)
def table_from_md(path):
    t = [l for l in path.read_text(encoding="utf-8").splitlines()]
    title = next((l[2:] for l in t if l.startswith("# ")), path.stem); doc.add_paragraph().add_run(title).bold = True
    i = 0
    while i < len(t):
        if t[i].startswith("|"):
            block = []
            while i < len(t) and t[i].startswith("|"): block.append(t[i]); i += 1
            rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in block if not re.match(r"^\|\s*:?-{3,}", r)]
            tb = doc.add_table(rows=len(rows), cols=len(rows[0])); tb.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, c in enumerate(row[:len(rows[0])]):
                    p = tb.cell(ri, ci).paragraphs[0]; r_ = p.add_run(c); r_.font.size = Pt(8.5); r_.bold = ri == 0
        else:
            if t[i].strip() and not t[i].startswith("# "):
                if t[i].startswith("## "): doc.add_paragraph().add_run(t[i][3:]).italic = True
                else: p = doc.add_paragraph(); runs(p, t[i]); [setattr(r.font, "size", Pt(9)) for r in p.runs]
            i += 1
    doc.add_paragraph()
in_math = False
for line in lines:
    if line.startswith("## Tables"): break
    if line.strip().startswith("$$"):
        p = doc.add_paragraph(); p.alignment = 1; runs(p, line.strip().strip("$")); continue
    if line.startswith("# "): h = doc.add_heading(line[2:], level=0)
    elif line.startswith("#### "): doc.add_heading(line[5:], level=3)
    elif line.startswith("### "): doc.add_heading(line[4:], level=2)
    elif line.startswith("## "): doc.add_heading(line[3:], level=1)
    elif line.startswith("- "): runs(doc.add_paragraph(style="List Bullet"), line[2:])
    elif line.strip() == "---": doc.add_page_break()
    elif line.strip(): runs(doc.add_paragraph(), line.strip())
doc.add_page_break(); doc.add_heading("Tables", level=1)
for f in ("Table1_planning_hierarchy.md", "Table2_bundle_domains.md", "Table3_bundle_completion.md", "Table4_report_cards.md"):
    table_from_md(R / "tables" / f)
doc.add_page_break(); doc.add_heading("Figures", level=1)
caps = {}
cap_block = (R / "manuscript_draft_en.md").read_text(encoding="utf-8").split("## Figure captions")[1].split("## Tables")[0]
for m in re.finditer(r"\*\*Fig\. (\d)\.\*\* (.+)", cap_block): caps[m.group(1)] = m.group(2)
for n, fn in enumerate(["Fig1_study_area.png", "Fig2_framework.png", "Fig3_completion_curves.png", "Fig4_single_facility_shortfall.png", "Fig5_living_zone_map_2025.png", "Fig6_zero_completion_by_unit.png"], 1):
    doc.add_picture(str(R / "figures" / fn), width=Cm(16)); p = doc.add_paragraph(); r_ = p.add_run(f"Fig. {n}. "); r_.bold = True; p.add_run(caps.get(str(n), ""))
out = R / "manuscript_draft_en.docx"; doc.save(out); print(out, out.stat().st_size)
