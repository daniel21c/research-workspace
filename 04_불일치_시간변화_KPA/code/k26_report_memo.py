# -*- coding: utf-8 -*-
"""
k26 — 교신저자 보고 메모(Word·PDF)를 Teams 보고글에서 만든다.

입력: 보고_20260930/Teams_보고글_20260930.md  (첫 '---' 줄 아래가 본문, 인사말 두 줄은 메모에서 뺀다)
출력: 보고_20260930/KPA_교수님보고_20260930.docx·.pdf   (PDF는 Word 자동화, k08.word_pdf)
실행: python k26_report_memo.py   (Windows + Word). 패키지 자체 점검 대상이 아니라 k20 목록에는 넣지 않는다.
"""
from __future__ import annotations
import re, sys
from pathlib import Path
import config as C

R = C.ROOT / "보고_20260930"
SRC = R / "Teams_보고글_20260930.md"
OUT = R / "KPA_교수님보고_20260930.docx"
TITLE = "연구4 KPA 투고 보고 메모 (2026-10-02)"
SUB = ("논문: 서울시 생활권의 내부통행률 상승과 경계 불일치의 지속 — 무작위 경계와 데이터 기반 커뮤니티를 이용한 경계 동 진단(2020·2025). "
       "「국토계획」 투고용, 심사용 16쪽.")
BASIS = ("설계: 연구설계.md(§3 IFR 분모 확정 근거) · 수치 대조: manuscript/수치대조_기록_v2.json(40/40) · 표 대조: manuscript/독립재계산_기록.json · "
         "IoU–D: results/t10_iou_vs_gap.csv · 재현성: results/_repro_result_benchmark.json(16/16) · 패키지 자체 점검: package/package_selfcheck.json · "
         "검수 대응: 검수의견_대응표_20260930.md(1~9차) · KPA 서식 근거: 검수기록/8차_KPA서식_레퍼런스점검_20261001.md, "
         "검수기록/9차_표선서식_분모확정_20261002.md")


def parse(md: str):
    body = md.split("\n---\n", 1)[1].strip().splitlines()
    items, k = [], 0
    for ln in body:
        s = ln.rstrip()
        if not s: continue
        if s.startswith("교수님, 안녕하세요") or s.startswith("연구4(국토계획 투고) 원고가"): continue
        if s.startswith("■ "): items.append(("h2", s[2:].strip())); continue
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m and not ln.startswith(" "): items.append(("num", (m.group(1), m.group(2)))); continue
        items.append(("p", s.strip()))
    return items


def build():
    import docx
    from docx.shared import Pt, Mm, RGBColor
    from docx.oxml.ns import qn
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(25); sec.top_margin = sec.bottom_margin = Mm(22)

    def font(style, size, bold=None, color=None):
        f = style.font; f.name = "맑은 고딕"; f.size = Pt(size)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        if bold is not None: f.bold = bold
        if color is not None: f.color.rgb = RGBColor(*color)
    font(d.styles["Normal"], 10.5)
    font(d.styles["Heading 1"], 16, True, (0x1F, 0x2A, 0x44)); font(d.styles["Heading 2"], 12.5, True, (0x1F, 0x2A, 0x44))
    nf = d.styles["Normal"].paragraph_format; nf.space_after = Pt(5); nf.line_spacing = 1.35

    d.add_heading(TITLE, level=1)
    d.add_paragraph(SUB)
    for kind, v in parse(SRC.read_text(encoding="utf-8")):
        if kind == "h2": d.add_heading(v, level=2)
        elif kind == "num":
            p = d.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.left_indent = Mm(6); p.paragraph_format.first_line_indent = Mm(-6)
            r = p.add_run(f"{v[0]}. {v[1]}"); r.bold = True
        else:
            p = d.add_paragraph(v)
            if p.text.startswith("첨부:"): p.paragraph_format.space_before = Pt(6)
    d.add_heading("확정 근거", level=2)
    d.add_paragraph(BASIS)
    d.save(str(OUT))
    return OUT


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.path.insert(0, str(Path(__file__).parent))
    from k08_hwp_pages import word_pdf
    out = build()
    pdf, n = word_pdf(out.resolve())
    print(f"memo: {out.name} → {pdf.name} {n}쪽")
