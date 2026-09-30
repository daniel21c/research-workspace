# -*- coding: utf-8 -*-
"""
k19 — 국토계획 투고 초본 v2(경계 동 진단·재배정) docx 생성과 본문 수치 자동 대조

원고: manuscript/국토계획_원고_<날짜>.md
표·그림: k18_v2_results.build_tables_v2(), FIGS_V2 (결과 파일에서 직접)
서식 도우미: k06_kpa_submission(국토계획 서식 — A4, HCR Batang 9.5pt, Ⅰ.1.1), 2단 편집본도 생성
수치 대조: 본문에 쓴 주장을 결과 파일과 대조해 manuscript/수치대조_기록_v2.json
실행: python k19_kpa_v2.py   (k13~k18 먼저)
"""
from __future__ import annotations
import json, re, sys
from pathlib import Path
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
import config as C
import k06_kpa_submission as K6
from k18_v2_results import build_tables_v2, FIGS_V2, draw_all, fig_framework

MK = C.MK; FIG = C.FIG; B = C.TAB / "benchmark"
WIDE_V2 = ("그림 1.", "그림 3.", "그림 5.", "표 1.", "표 2.", "표 3.", "표 4.", "표 5.", "표 6.", "표 A1.", "표 A2.", "그림 A1.")


def cap(doc, ko, en, kind):
    """편집규정 제19조⑨⑩: 표 제목은 표 위 왼쪽, 그림 제목은 그림 아래 가운데, 국문 줄 + 영문 줄."""
    p = doc.add_paragraph(style="KPA Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if kind == "table" else WD_ALIGN_PARAGRAPH.CENTER
    for k, (lab, txt) in enumerate(((("표" if kind == "table" else "그림"), ko), (("Table" if kind == "table" else "Fig."), en))):
        num, rest = txt.split(" ", 1)
        r = p.add_run(f"{lab} {num} "); K6.run_font(r, K6.BODY_FONT, 8.5, bold=True)
        r = p.add_run(rest); K6.run_font(r, K6.BODY_FONT, 8.5)
        if k == 0: r.add_break()
    p.paragraph_format.keep_with_next = kind == "table"
    return p


def build(md_path: Path):
    meta, body = {}, []
    for line in md_path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^%([A-Z_]+):\s*(.*)$", line)
        if m: meta[m.group(1)] = m.group(2).strip()
        else: body.append(line)
    tables = build_tables_v2()
    doc = Document()
    sec = doc.sections[0]; sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(1.1); sec.left_margin = Cm(2.0); sec.right_margin = Cm(1.8)
    S = K6.style; A = WD_ALIGN_PARAGRAPH
    S(doc, "Normal", K6.BODY_FONT, 9.5, line=1.6, first=0.35)
    # 편집규정 제19조: 제목·부제 좌측 정렬, 부제는 ':'로 시작해 행 분리, 요약문 양쪽 정렬
    S(doc, "KPA Title", K6.HEAD_FONT, 14, bold=True, align=A.LEFT, line=1.3, after=2)
    S(doc, "KPA Subtitle", K6.HEAD_FONT, 11, align=A.LEFT, line=1.3, after=6)
    S(doc, "KPA Title EN", K6.EN_FONT, 12, bold=True, align=A.LEFT, line=1.2, before=4, after=2)
    S(doc, "KPA Subtitle EN", K6.EN_FONT, 10.5, align=A.LEFT, line=1.2, after=10)
    S(doc, "KPA Abstract", K6.EN_FONT, 8.5, align=A.JUSTIFY, line=1.3, after=6)
    S(doc, "KPA Keywords", K6.BODY_FONT, 8.5, align=A.LEFT, line=1.3, after=2)
    S(doc, "KPA Heading 1", K6.HEAD_FONT, 13.5, bold=True, align=A.LEFT, line=1.2, before=14, after=6)
    S(doc, "KPA Heading 2", K6.HEAD_FONT, 11, bold=True, align=A.LEFT, line=1.2, before=10, after=4)
    S(doc, "KPA Caption", K6.BODY_FONT, 8.5, bold=True, align=A.CENTER, line=1.1, before=6, after=3)
    S(doc, "KPA Note", K6.BODY_FONT, 7.5, line=1.15, after=8)
    S(doc, "KPA Reference", K6.BODY_FONT, 8.5, align=A.LEFT, line=1.3, after=2)
    doc.styles["KPA Reference"].paragraph_format.left_indent = Cm(0.8); doc.styles["KPA Reference"].paragraph_format.first_line_indent = Cm(-0.8)
    K6.para(doc, meta["TITLE_KO"], "KPA Title"); K6.para(doc, ": " + meta["SUBTITLE_KO"], "KPA Subtitle")
    K6.para(doc, meta["TITLE_EN"], "KPA Title EN"); K6.para(doc, ": " + meta["SUBTITLE_EN"], "KPA Subtitle EN")
    p = doc.add_paragraph(style="KPA Abstract"); r = p.add_run("Abstract"); K6.run_font(r, K6.EN_FONT, 9, bold=True)
    p = doc.add_paragraph(style="KPA Abstract"); r = p.add_run(meta["ABSTRACT_EN"]); K6.run_font(r, K6.EN_FONT, 8.5)
    for lab, key, font in (("주제어: ", "KEYWORDS_KO", K6.BODY_FONT), ("Keywords: ", "KEYWORDS_EN", K6.EN_FONT)):
        p = doc.add_paragraph(style="KPA Keywords"); r = p.add_run(lab); K6.run_font(r, font, 8.5, bold=True); r = p.add_run(meta[key]); K6.run_font(r, font, 8.5)
    in_refs = False; ref_no = 0
    for line in body:
        s = line.strip()
        if not s: continue
        if s.startswith("# "): t = s[2:].strip(); in_refs = t.startswith("인용문헌"); K6.para(doc, t, "KPA Heading 1"); continue
        if s.startswith("## "): K6.para(doc, s[3:], "KPA Heading 2"); continue
        m = re.match(r"^\[\[TABLE:(\w+)\]\]$", s)
        if m:
            t = tables[m.group(1)]; cap(doc, t["ko"], t["en"], "table")
            K6.add_table(doc, t["headers"], t["rows"], t["widths"], bold_last=t.get("bold_last", False)); K6.note(doc, t["note"]); continue
        m = re.match(r"^\[\[FIG:(\w+)\]\]$", s)
        if m:
            ko, en, fn, w, nt = FIGS_V2[m.group(1)]; p = doc.add_paragraph(); p.alignment = A.CENTER
            p.add_run().add_picture(str(FIG / fn), width=Cm(w)); cap(doc, ko, en, "figure"); K6.note(doc, nt); continue
        if in_refs:                                  # 번호 + 국문 문헌, 줄바꿈 후 영문 병기(학회 샘플 형식)
            ref_no += 1; parts = s.split(" // "); p = doc.add_paragraph(style="KPA Reference")
            for k, t in enumerate(parts):
                r = p.add_run((f"{ref_no}. " if k == 0 else "") + t); K6.run_font(r, K6.BODY_FONT, 8.5)
                if k < len(parts) - 1: r.add_break()
            continue
        K6.para(doc, s, "Normal")
    cp = doc.core_properties; cp.author = ""; cp.last_modified_by = ""; cp.title = ""; cp.comments = ""
    out = C.DOCX / (md_path.stem + ".docx"); doc.save(out)
    K6.WIDE_CAPTIONS = WIDE_V2
    out2 = K6.two_column(out, out.with_name(out.stem + "_2단편집.docx"))
    return out, out2


def claims(md_text: str):
    """(1) 원고 본문 수치 대조(k24: 원고에 적힌 숫자 ↔ 결과 파일, 절·문맥 단위) + 변조 시험, (2) 구조 점검, (3) 인용 ↔ 참고문헌."""
    import k24_text_claims as K24
    rc = lambda n: pd.read_csv(n, encoding="utf-8-sig"); b = lambda n: json.loads((B / n).read_text(encoding="utf-8"))
    S4, ver = b("b4_summary.json"), b("b8_verify.json"); gs = rc(B / "b4_gu_summary.csv")
    V = K24.values(); text = K24.check(md_text, V); mut = K24.mutation_test(md_text, V)
    Ck = []
    def ck(n, c, detail=""): Ck.append({"항목": n, "일치": bool(c), "비고": detail})
    tabs = build_tables_v2()
    ck("재배정 분할 검증(권역 수·연속·요약값)", ver["판정"] == "통과", json.dumps(ver, ensure_ascii=False)[:200])
    ck("표 A1 행 수 = 원래→최종 생활권이 두 해 같은 동 수", len(tabs["TA1"]["rows"]) == S4["두해모두_권고_이동"])
    gj = gs[(gs.year == 2025) & (gs.ku_name == "광진구")].iloc[0]; ck("광진 재배정 후 커뮤니티와 거의 같음(D_after < 0.5%)", gj.D_after < 0.005, f"D_after={gj.D_after:.4f}")
    t3 = tabs["T3"]; S9 = b("b9_change_story.json")
    ck("표 3 유형 내 효과 비중 > 50%(사전 기준 C1) = 본문", S9["유형분해"]["유형내_비중"] > 0.5 and S9["유형분해"]["판정(유형내>50%)"] == "통과")
    ck("주말 국지화(사전 기준 C3) 통과", S9["평일주말"]["판정"] == "통과")
    ck("표 4 행 = 불일치·격차·분해 4행", len(tabs["T4"]["rows"]) == 4)
    ab = {k: re.search(rf"^%{k}: (.*)$", md_text, flags=re.M).group(1) for k in ("ABSTRACT_EN", "ABSTRACT_KO")}
    nums = {k: [x for x in re.findall(r"\d+(?:[.,]\d+)?", t) if x not in ("2020", "2025")] for k, t in ab.items()}
    ck("초록에 구체적 수치 없음(연도만 허용)", not any(nums.values()), json.dumps(nums, ensure_ascii=False))
    ck("원고에 원인 탐색·시설 접근성 표(구 표 6) 없음", "[[TABLE:T6]]" in md_text and "Coverage" not in md_text)
    bad = [w for w in ("가까운 곳에서 이루어", "가까운 곳으로 몰린", "이들 동만", "공간적으로 연결된 커뮤니티", "크기 효과를 통제", "only those dongs", "불일치는 줄지 않았", "따라가지 못하", "8~9%에 해당하는 일부 경계 동에 집중", "쓸모없는 경계", "저절로 해소", "더 멀어져", "Claude", "ChatGPT", "박사") if w in md_text]
    ck("계산이 뒷받침하지 않는 표현 없음(AI 대조 점검 1·2차 09-30: 거리 단축·진단 동만 재배정·Leiden 공간 연결·크기 통제·D 불변·8~9%에 D 집중·따라가지 못함·전면 재설정; 6차 저자 결정으로 '집중'·'잘 그어짐'은 허용) + 저자 결정(도구명·학위논문 언급 없음; 규정 적합성 판정이 아님)", not bad, str(bad))
    ai = md_text.split("# AI 사용 진술문")[1].split("# 인용문헌")[0]
    ck("AI 진술문에 사용 목적·시기·저자 검토 문장 있음(가이드라인 가.5; 도구명은 저자 결정으로 미기재 — 체크리스트 B-7)", all(w in ai for w in ("2026년 9월", "검토·수정·검증", "책임")), ai[:80])
    words = len(re.search(r"^%ABSTRACT_EN: (.*)$", md_text, flags=re.M).group(1).split())
    ck("영문 초록 200단어 내외(편집규정 제19조③)", 180 <= words <= 230, f"{words}단어")
    ck("원고 본문 수치 대조(k24) 전부 일치", all(c["일치"] for c in text), str([c["ID"] for c in text if not c["일치"]]))
    ck("변조 시험: 숫자를 바꾸거나 문장을 지우면 모두 실패로 감지", not mut["감지못함"] and mut["숫자변조_감지"] == mut["숫자변조"] > 0 and mut["삭제_감지"] == mut["삭제"] > 0, json.dumps(mut, ensure_ascii=False))
    # 인용 ↔ 참고문헌
    body, refs = md_text.split("# 인용문헌")[0].split("# AI 사용 진술문")[0], md_text.split("# 인용문헌")[1].split("# 부록")[0]
    keys = set()
    for c in re.findall(r"\(([^()]*?(?:19|20)\d{2}[^()]*?)\)", body):
        for part in c.split(";"):
            m = re.search(r"([A-Za-z가-힣][A-Za-z가-힣·\s\-\.]*?)(?:\s*et al\.|\s*외)?,?\s*((?:19|20)\d{2})(?!년)", part.strip())
            if m: keys.add((m.group(1).strip(), m.group(2)))
    lines = [l for l in refs.splitlines() if l.strip()]
    miss = [f"{a} {y}" for a, y in keys if not any(y in l and a.split(" and ")[0].split()[0][:3] in l for l in lines)]
    unc = [l[:30] for l in lines if not any(y in l and a.split(" and ")[0].split()[0][:3] in l for a, y in keys) and "법률" not in l]
    ck("인용→참고문헌 누락 없음", not miss, str(miss)); ck("참고문헌→인용 미인용 없음", not unc, str(unc))
    out = {"설명": "원고 수치 대조는 원고(md)의 절별 문장에서 결과 파일 값(원정밀도 → 표시 반올림 한 번)을 찾는다(k24). 변조 시험은 주장마다 숫자를 바꾸거나 문장을 지운 원고로 다시 검사해 실패하는지 본다.",
           "원고_수치_대조": {"주장수": len(text), "일치": sum(c["일치"] for c in text), "항목": text}, "변조_시험": mut,
           "항목": Ck, "실패": [c["항목"] for c in Ck if not c["일치"]]}
    (MK / "수치대조_기록_v2.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    return out


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    mds = sorted(MK.glob("국토계획_원고_*.md"))
    if not mds: raise SystemExit(f"원고 md가 없다: {MK}/국토계획_원고_<날짜>.md (패키지는 manuscript/ 폴더)")
    md = mds[-1]
    draw_all()
    o1, o2 = build(md); print("docx:", o1.name, "|", o2.name)
    r = claims(md.read_text(encoding="utf-8"))
    for c in r["항목"]: print("OK  " if c["일치"] else "FAIL", c["항목"], "" if c["일치"] else c["비고"])
    print(f"원고 수치 대조 {r['원고_수치_대조']['일치']}/{r['원고_수치_대조']['주장수']}, 변조 시험 {r['변조_시험']}")
    print("총", len(r["항목"]), "실패", len(r["실패"]))
