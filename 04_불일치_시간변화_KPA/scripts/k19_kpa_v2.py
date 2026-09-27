# -*- coding: utf-8 -*-
"""
k19 — 국토계획 투고 초본 v2(경계 동 진단·재배정) docx 생성과 본문 수치 자동 대조

원고: output/manuscript_kpa/국토계획_투고초본_v2_<날짜>.md
표·그림: k18_v2_results.build_tables_v2(), FIGS_V2 (결과 파일에서 직접)
서식 도우미: k06_kpa_submission(국토계획 서식 — A4, HCR Batang 9.5pt, Ⅰ.1.1), 2단 편집본도 생성
수치 대조: 본문에 쓴 주장을 결과 파일과 대조해 output/manuscript_kpa/수치대조_기록_v2.json
실행: python k19_kpa_v2.py   (k13~k18 먼저)
"""
from __future__ import annotations
import json, re, sys
from pathlib import Path
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
import config as C
import k06_kpa_submission as K6
from k18_v2_results import build_tables_v2, FIGS_V2, draw_all, fig_framework

MK = C.OUT / "manuscript_kpa"; FIG = C.OUT / "figures"; B = C.TAB / "benchmark"
WIDE_V2 = ("그림 1.", "그림 3.", "그림 5.", "표 1.", "표 3.", "표 4.", "표 5.", "표 6.", "표 A1.", "그림 A1.")


def build(md_path: Path):
    meta, body = {}, []
    for line in md_path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^%([A-Z_]+):\s*(.*)$", line)
        if m: meta[m.group(1)] = m.group(2).strip()
        else: body.append(line)
    tables = build_tables_v2()
    doc = Document()
    sec = doc.sections[0]; sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(1.1); sec.left_margin = Cm(2.0); sec.right_margin = Cm(1.8)
    S = K6.style; A = WD_ALIGN_PARAGRAPH
    S(doc, "Normal", K6.BODY_FONT, 9.5, line=1.6, first=0.35)
    S(doc, "KPA Title", K6.HEAD_FONT, 14, bold=True, align=A.CENTER, line=1.3, after=2)
    S(doc, "KPA Subtitle", K6.HEAD_FONT, 11, align=A.CENTER, line=1.3, after=6)
    S(doc, "KPA Title EN", K6.EN_FONT, 12, bold=True, align=A.CENTER, line=1.2, before=4, after=2)
    S(doc, "KPA Subtitle EN", K6.EN_FONT, 10.5, align=A.CENTER, line=1.2, after=10)
    S(doc, "KPA Abstract", K6.EN_FONT, 8.5, line=1.3, after=6)
    S(doc, "KPA Keywords", K6.BODY_FONT, 8.5, align=A.LEFT, line=1.3, after=2)
    S(doc, "KPA Heading 1", K6.HEAD_FONT, 13.5, bold=True, align=A.LEFT, line=1.2, before=14, after=6)
    S(doc, "KPA Heading 2", K6.HEAD_FONT, 11, bold=True, align=A.LEFT, line=1.2, before=10, after=4)
    S(doc, "KPA Caption", K6.BODY_FONT, 8.5, bold=True, align=A.CENTER, line=1.1, before=6, after=3)
    S(doc, "KPA Note", K6.BODY_FONT, 7.5, line=1.15, after=8)
    S(doc, "KPA Reference", K6.BODY_FONT, 8.5, align=A.LEFT, line=1.3, after=2)
    doc.styles["KPA Reference"].paragraph_format.left_indent = Cm(0.8); doc.styles["KPA Reference"].paragraph_format.first_line_indent = Cm(-0.8)
    K6.para(doc, meta["TITLE_KO"], "KPA Title"); K6.para(doc, "- " + meta["SUBTITLE_KO"] + " -", "KPA Subtitle")
    K6.para(doc, meta["TITLE_EN"], "KPA Title EN"); K6.para(doc, "- " + meta["SUBTITLE_EN"] + " -", "KPA Subtitle EN")
    p = doc.add_paragraph(style="KPA Abstract"); r = p.add_run("Abstract  "); K6.run_font(r, K6.EN_FONT, 9, bold=True); r = p.add_run(meta["ABSTRACT_EN"]); K6.run_font(r, K6.EN_FONT, 8.5)
    for lab, key, font in (("주제어: ", "KEYWORDS_KO", K6.BODY_FONT), ("Keywords: ", "KEYWORDS_EN", K6.EN_FONT)):
        p = doc.add_paragraph(style="KPA Keywords"); r = p.add_run(lab); K6.run_font(r, font, 8.5, bold=True); r = p.add_run(meta[key]); K6.run_font(r, font, 8.5)
    in_refs = False
    for line in body:
        s = line.strip()
        if not s: continue
        if s.startswith("# "): t = s[2:].strip(); in_refs = t.startswith("인용문헌"); K6.para(doc, t, "KPA Heading 1"); continue
        if s.startswith("## "): K6.para(doc, s[3:], "KPA Heading 2"); continue
        m = re.match(r"^\[\[TABLE:(\w+)\]\]$", s)
        if m:
            t = tables[m.group(1)]; K6.caption(doc, t["ko"], t["en"], "table")
            K6.add_table(doc, t["headers"], t["rows"], t["widths"], bold_last=t.get("bold_last", False)); K6.note(doc, t["note"]); continue
        m = re.match(r"^\[\[FIG:(\w+)\]\]$", s)
        if m:
            ko, en, fn, w, nt = FIGS_V2[m.group(1)]; p = doc.add_paragraph(); p.alignment = A.CENTER
            p.add_run().add_picture(str(FIG / fn), width=Cm(w)); K6.caption(doc, ko, en, "figure"); K6.note(doc, nt); continue
        K6.para(doc, s, "KPA Reference" if in_refs else "Normal")
    cp = doc.core_properties; cp.author = ""; cp.last_modified_by = ""; cp.title = ""; cp.comments = ""
    out = MK / (md_path.stem + ".docx"); doc.save(out)
    K6.WIDE_CAPTIONS = WIDE_V2
    out2 = K6.two_column(out, out.with_name(out.stem + "_2단편집.docx"))
    return out, out2


def claims(md_text: str):
    rc = lambda n: pd.read_csv(n, encoding="utf-8-sig")
    res = json.loads((C.TAB / "results.json").read_text(encoding="utf-8")); b = lambda n: json.loads((B / n).read_text(encoding="utf-8"))
    S1, S4, S5, S6, S7 = b("b_summary.json"), b("b4_summary.json"), b("b5_summary.json"), b("b6_summary.json"), b("b7_summary.json")
    g = rc(B / "b1_gu.csv"); z = rc(B / "b2_zone.csv"); gs = rc(B / "b4_gu_summary.csv"); gm = rc(B / "b4_greedy_moves.csv"); ver = b("b8_verify.json")
    sel = rc(C.TAB / "t06_selection.csv"); d3 = rc(B / "b3_dong.csv")
    r1 = lambda x: round(x * 100, 1); r0 = lambda x: round(x * 100)
    Ck = []
    def ck(n, c, detail=""): Ck.append({"항목": n, "일치": bool(c), "비고": detail})
    Sd = res["seoul"]
    ck("IFR LZ 34.1→37.9, LD 34.8→39.4", (r1(Sd["2020"]["IFR_lz"]), r1(Sd["2025"]["IFR_lz"]), r1(Sd["2020"]["IFR_ld"]), r1(Sd["2025"]["IFR_ld"])) == (34.1, 37.9, 34.8, 39.4))
    ck("무작위 범위 안 23/25", res["null"]["lz_within_null90_n"] == 23); ck("동 내부통행 18.4→21.1", (r1(Sd["2020"]["SR"]), r1(Sd["2025"]["SR"])) == (18.4, 21.1))
    lzg = g[g.boundary == "LZ"]; med = lzg.groupby("year")[["IFR", "med_N0", "med_N1"]].median()
    ck("IFR 중앙값 N0 0.353 vs N1 0.332 (2025)", round(med.loc[2025, "med_N0"], 3) == 0.353 and round(med.loc[2025, "med_N1"], 3) == 0.332)
    ck("LZ 백분위 N0 55/57, N1 97/97, N2 97/98", (r0(S1["구_LZ_2020"]["pct_N0_중앙"]), r0(S1["구_LZ_2025"]["pct_N0_중앙"]), r0(S1["구_LZ_2020"]["pct_N1_중앙"]), r0(S1["구_LZ_2025"]["pct_N1_중앙"]), r0(S1["구_LZ_2020"]["pct_N2_중앙"]), r0(S1["구_LZ_2025"]["pct_N2_중앙"])) == (55, 57, 97, 97, 97, 98))
    ck("LD 백분위 N1 97/99", (r0(S1["구_LD_2020"]["pct_N1_중앙"]), r0(S1["구_LD_2025"]["pct_N1_중앙"])) == (97, 99))
    pv = lzg.pivot(index="ku_name", columns="year", values="pct_N1")
    ck("은평 17/35, 강동 88→51", (r0(pv.loc["은평구", 2020]), r0(pv.loc["은평구", 2025]), r0(pv.loc["강동구", 2020]), r0(pv.loc["강동구", 2025])) == (17, 35, 88, 51))
    ck("생활권 75/78%, 무작위보다 못한 29/25", (r0(S1["생활권_LZ_2020"]["무작위보다_나은_비율(pct>0.5)"]), r0(S1["생활권_LZ_2025"]["무작위보다_나은_비율(pct>0.5)"])) == (75, 78) and int(((z.boundary == "LZ") & (z.year == 2020) & (z.pct < 0.5)).sum()) == 29 and int(((z.boundary == "LZ") & (z.year == 2025) & (z.pct < 0.5)).sum()) == 25)
    lw = z[z.boundary == "LZ"].pivot(index="zone", columns="year", values="pct"); six = sorted(lw[(lw[2020] < 0.25) & (lw[2025] < 0.25)].index)
    ck("하위 25% 6개 생활권", {x.split("_", 1)[1].replace("생활권", "") for x in six} == {"개포일원", "쌍문", "충정", "후암용산", "수색", "청운효자"}, str(six))
    ck("옆 생활권 지향 135/119", (S1["동_LZ_2020"]["이웃권역이_더_담는_동수"], S1["동_LZ_2025"]["이웃권역이_더_담는_동수"]) == (135, 119))
    ck("진짜 오배정 35/40, 그 밖 0/0", (S4["2020_옆생활권지향동"]["둘다>0"], S4["2025_옆생활권지향동"]["둘다>0"], S4["2020_그밖의_경계동"]["둘다>0"], S4["2025_그밖의_경계동"]["둘다>0"]) == (35, 40, 0, 0))
    q5, q0 = S4["2025_탐욕재배정"], S4["2020_탐욕재배정"]
    ck("2025 65개·20구, IFR 37.9→39.6(LD 39.4), D 8.3→3.2, Q 0.387→0.395(LD 0.404)", (q5["옮긴_동"], q5["옮긴_구"], round(q5["서울IFR_전"], 1), round(q5["서울IFR_후"], 1), round(q5["서울IFR_가상경계"], 1), round(q5["서울D_전"], 1), float(f"{q5['서울D_후']:.1f}"), round(q5["Q_전_중앙"], 3), round(q5["Q_후_중앙"], 3), round(q5["Q_가상경계_중앙"], 3)) == (65, 20, 37.9, 39.6, 39.4, 8.3, 3.2, 0.387, 0.395, 0.404))
    ck("D 감소 60% 넘게", (q5["서울D_전"] - q5["서울D_후"]) / q5["서울D_전"] > 0.6)
    ck("2020 66개, 34.1→35.3(LD 34.8), D 8.6→3.7", (q0["옮긴_동"], round(q0["서울IFR_전"], 1), round(q0["서울IFR_후"], 1), round(q0["서울IFR_가상경계"], 1), round(q0["서울D_전"], 1), round(q0["서울D_후"], 1)) == (66, 34.1, 35.3, 34.8, 8.6, 3.7))
    ck("ΔIFR<0 이동 19/18", (int((gm[gm.year == 2020].dIFR_pp < 0).sum()), int((gm[gm.year == 2025].dIFR_pp < 0).sum())) == (19, 18))
    gj = gs[(gs.year == 2025) & (gs.ku_name == "광진구")].iloc[0]; ck("광진 6개·+2.9%p", int(gj.n_moved) == 6 and round(gj.dIFR_pp, 1) == 2.9)
    ck("광진 재배정 후 가상경계와 거의 같음(D_after < 0.5%)", gj.D_after < 0.005, f"D_after={gj.D_after:.4f}")
    ck("두 해 공통 49, 2020만 17, 2025만 16", (S4["두해모두_권고_이동"], S4["2020만"], S4["2025만"]) == (49, 17, 16))
    ck("LD2020 고정 시 D 증가 15", int(sel.D_increase_fixed_ld2020.sum()) == 15)
    ck("원인: 기초시설 관계없음(p 0.60/0.13)", round(S5["경계비용_COV_main"]["2020"]["p_A_vs_B"], 2) == 0.60 and round(S5["경계비용_COV_main"]["2025"]["p_A_vs_B"], 2) == 0.13)
    e = S6["결과"]; ck("역 +12%p 안팎, p 0.05/0.06", 11 <= e["2020"]["역"]["차이_%p"] <= 13 and 11 <= e["2025"]["역"]["차이_%p"] <= 13 and round(e["2020"]["역"]["p"], 2) == 0.05 and round(e["2025"]["역"]["p"], 2) == 0.06)
    ck("대형상업·문화 관계없음", S6["판정"]["대형상업"] == "지지 안 됨" and S6["판정"]["문화"] == "지지 안 됨")
    w7 = S7["결과"]; u = "업무성차_자기권역−동_중앙(배, exp)"; ck("중심 권역 1.2배 안팎, p 0.051/0.018", 1.15 <= w7["2020"][u]["A"] <= 1.25 and 1.15 <= w7["2025"][u]["A"] <= 1.25 and round(w7["2020"][u]["p"], 3) == 0.051 and round(w7["2025"][u]["p"], 3) == 0.018)
    ck("주거 특성 관계없음", S7["판정_유사성"] == "지지 안 됨")
    ck("재배정 분할 검증(권역 수·연속·요약값)", ver["판정"] == "통과", json.dumps(ver, ensure_ascii=False)[:200])
    # 인용 ↔ 참고문헌
    body, refs = md_text.split("# 인용문헌")[0], md_text.split("# 인용문헌")[1].split("# 부록")[0]
    keys = set()
    for c in re.findall(r"\(([^()]*?(?:19|20)\d{2}[^()]*?)\)", body):
        for part in c.split(";"):
            m = re.search(r"([A-Za-z가-힣][A-Za-z가-힣·\s\-\.]*?)(?:\s*et al\.|\s*외)?,?\s*((?:19|20)\d{2})(?!년)", part.strip())
            if m: keys.add((m.group(1).strip(), m.group(2)))
    lines = [l for l in refs.splitlines() if l.strip()]
    miss = [f"{a} {y}" for a, y in keys if not any(y in l and a.split(" and ")[0].split()[0][:3] in l for l in lines)]
    unc = [l[:30] for l in lines if not any(y in l and a.split(" and ")[0].split()[0][:3] in l for a, y in keys) and "법률" not in l]
    ck("인용→참고문헌 누락 없음", not miss, str(miss)); ck("참고문헌→인용 미인용 없음", not unc, str(unc))
    out = {"항목": Ck, "실패": [c["항목"] for c in Ck if not c["일치"]]}
    (MK / "수치대조_기록_v2.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    return out


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    md = sorted(MK.glob("국토계획_투고초본_v2_*.md"))[-1]
    draw_all()
    o1, o2 = build(md); print("docx:", o1.name, "|", o2.name)
    r = claims(md.read_text(encoding="utf-8"))
    for c in r["항목"]: print("OK  " if c["일치"] else "FAIL", c["항목"], "" if c["일치"] else c["비고"])
    print("총", len(r["항목"]), "실패", len(r["실패"]))
