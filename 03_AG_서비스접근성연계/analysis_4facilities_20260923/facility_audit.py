"""Offline full-record lineage checks; never imports or runs the original builder."""
import ast
import hashlib
from collections import Counter
import io
import json
from pathlib import Path
import zipfile
import numpy as np
import pandas as pd
import geopandas as gpd
import openpyxl
from pyproj import Transformer
import config as C

def decoded(name):
    try:return name.encode('cp437').decode('cp949')
    except (UnicodeError,LookupError):return name

def retail_members(z,part,prefix=''):
    for info in z.infolist():
        name=decoded(info.filename)
        if name.endswith('.zip'):
            with zipfile.ZipFile(io.BytesIO(z.read(info))) as inner:
                yield from retail_members(inner,part,prefix+name+'!')
        elif part in name and name.endswith('.csv'):
            with z.open(info) as f:
                yield prefix+name,pd.read_csv(f,dtype=str,encoding='utf-8-sig',keep_default_na=False)

def school_roster(year):
    if year==2020:
        p=C.RAW_FACILITIES/'school/2019년 하반기 유초중고 학교 주소록_20191001.xlsx';ws='Sheet1';skip=4;cols=(4,1,7,13,10,8,9)
    else:
        p=C.RAW_FACILITIES/'school/2024년 하반기 교육통계 학교별 일람표.xlsx';ws='학교별 주요 통계';skip=13;cols=(1,4,8,19,11,10,14)
    w=openpyxl.load_workbook(p,read_only=True,data_only=True);records=[]
    for i,r in enumerate(w[ws].iter_rows(values_only=True)):
        if i<skip:continue
        region,typ,name,address,est,branch,state=[r[j] for j in cols]
        if region=='서울' and typ in ('초등학교','중학교','고등학교') and state in ('기존(원)교','신설(원)교') and (year==2020 or branch=='본교'):
            records.append({'facility_name':str(name).strip(),'road_address':str(address).strip() if address else '', 'establishment':est,'branch':branch,'raw_school_type':typ})
    w.close();return p,pd.DataFrame(records)

def library_roster(year):
    p=C.RAW_FACILITIES/f'library/library_public_{year-1}_reference.xlsx';w=openpyxl.load_workbook(p,read_only=True,data_only=True);records=[]
    for i,r in enumerate(w[w.sheetnames[0]].iter_rows(values_only=True)):
        if i<(1 if year==2020 else 2):continue
        if r[3]=='서울' and r[2] in ('공공(일반)','공공(어린이)'):
            a=r[12 if year==2020 else 11]
            records.append({'facility_name':str(r[1]).strip(),'road_address':str(a).strip() if a else '', 'raw_library_type':r[2]})
    w.close();return p,pd.DataFrame(records)

def audit():
    inputs=set();summary=[];all_rows=[];counts=[]
    evidence_records=[]
    for base in C.RETAIL_SOURCE_VERIFICATIONS:
        evidence=dict(base);ey=evidence['year_label']
        filename='retail_2020_FILE_000000003547804.zip' if ey==2020 else 'retail_2025_FILE_000000003676580.zip'
        official_local=C.RAW_FACILITIES/'retail'/filename
        with official_local.open('rb') as stream:h=hashlib.file_digest(stream,'sha256').hexdigest()
        evidence['local_sha256']=h;evidence['local_bytes']=official_local.stat().st_size
        assert h==evidence['sha256'] and evidence['local_bytes']==evidence['bytes']
        evidence_records.append(evidence)
        summary.append({'year':ey,'category':'RETAIL_DAILY','check':'current_official_download_payload','status':'PASS','count':evidence['bytes'],'detail':f'Coordinator verified current official HTTPS payload on {evidence["verification_date"]}; exact local byte size and SHA256 match. Historical content vintage is a separate question.'})
    (C.RESULTS/'external_source_verification.json').write_text(json.dumps(evidence_records,ensure_ascii=False,indent=2),encoding='utf-8')
    builder=C.FACILITIES/'build_seoul_4facilities_clean_2020_2025.py';inputs.add(builder)
    # Extract a literal codebook only; original module has credentials/API side effects.
    tree=ast.parse(builder.read_text('utf-8-sig'));codebook={}
    for node in tree.body:
        if isinstance(node,(ast.Assign,ast.AnnAssign)):
            targets=node.targets if isinstance(node,ast.Assign) else [node.target]
            if any(isinstance(t,ast.Name) and t.id=='RETAIL_SUBTYPES' for t in targets): codebook=ast.literal_eval(node.value)
    assert codebook
    gu=gpd.read_file(C.GU).to_crs(4326).geometry.union_all()
    inputs.update(C.GU.parent.glob(C.GU.stem+'.*'))
    clinic_path=C.RAW_FACILITIES/'hospital/seoul_clinic_health_permits_current_with_closed_OA-16480.csv';inputs.add(clinic_path)
    clinic=pd.read_csv(clinic_path,encoding='cp949',dtype=str,keep_default_na=False).set_index('관리번호')
    geocode_path=C.FACILITIES/'geocode_input_address.xlsx';inputs.add(geocode_path)
    geocode=pd.read_excel(geocode_path).set_index(['year_snapshot','facility_id'])
    trans=Transformer.from_crs(5174,4326,always_xy=True)
    for y in C.YEARS:
        cp=C.FACILITIES/f'seoul_facilities_{y}_01_clean.csv';pp=cp.with_suffix('.parquet');inputs.update([cp,pp])
        c=pd.read_csv(cp,keep_default_na=False).sort_values('facility_id').reset_index(drop=True);p=pd.read_parquet(pp).sort_values('facility_id').reset_index(drop=True)
        strings=[x for x in c.columns if x not in ('longitude','latitude')]
        string_equal=bool(c[strings].fillna('').astype(str).equals(p[strings].fillna('').astype(str)))
        coord_error=float(np.abs(c[['longitude','latitude']].to_numpy()-p[['longitude','latitude']].to_numpy()).max())
        assert len(c)==len(p) and string_equal and coord_error<1e-12
        assert c.facility_id.is_unique
        geo=gpd.GeoSeries(gpd.points_from_xy(c.longitude,c.latitude),crs=4326)
        inside=geo.within(gu);covered=geo.covered_by(gu)
        row=c[['facility_id','facility_category','facility_name','year_snapshot','source_name']].copy()
        row['year']=y;row['inside_seoul_polygon']=inside;row['on_seoul_boundary']=covered&~inside
        row['source_row_match']=False;row['source_coordinate_match']=False;row['coordinate_evidence_status']='UNVERIFIED';row['establishment']='';row['historical_coordinate_verified']=False
        summary.append({'year':y,'category':'ALL','check':'csv_parquet_all_values','status':'PASS','count':len(c),'detail':f'All strings equal; max coordinate representation difference {coord_error:.3g}'})
        summary.append({'year':y,'category':'ALL','check':'actual_seoul_polygon','status':'PASS' if covered.all() else 'FAIL','count':int(covered.sum()),'detail':f'{len(c)} rows, strict inside={inside.sum()}, boundary={sum(covered&~inside)}, outside={sum(~covered)}'})
        counts.extend(c.groupby(['facility_category','facility_sub']).size().reset_index(name='count').assign(year=y).to_dict('records'))
        for cat in C.CATEGORIES:
            mask=c.facility_category.eq(cat);g=c.loc[mask]
            if cat=='RETAIL_DAILY':
                rawp=C.RAW_FACILITIES/('retail/retail_2020_FILE_000000003547804.zip' if y==2020 else 'retail/retail_2025_FILE_000000003676580.zip');inputs.add(rawp)
                with zipfile.ZipFile(rawp) as z:members=list(retail_members(z,'서울_201912' if y==2020 else '서울_202412'))
                assert len(members)==1
                member,raw=members[0];raw['facility_id']='RETAIL_'+raw['상가업소번호'].str.strip();raw=raw.set_index('facility_id')
                mapped=raw.reindex(g.facility_id)
                valid=mapped['상호명'].fillna('').str.strip().to_numpy()==g.facility_name.to_numpy()
                valid&=mapped['도로명주소'].fillna('').str.strip().to_numpy()==g.road_address.to_numpy()
                valid&=mapped['상권업종소분류코드'].map(codebook).to_numpy()==g.facility_sub.to_numpy()
                coords=np.all(np.isclose(mapped[['경도','위도']].apply(pd.to_numeric,errors='coerce'),g[['longitude','latitude']],atol=1e-8,rtol=0),axis=1)
                row.loc[mask,'source_row_match']=valid;row.loc[mask,'source_coordinate_match']=coords
                row.loc[mask,'coordinate_evidence_status']=np.where(coords,'MATCH_PRESERVED_RAW','REVIEW_MISMATCH')
                summary.append({'year':y,'category':cat,'check':'full_raw_row_link','status':'PASS' if valid.all() and coords.all() else 'REVIEW','count':int((valid&coords).sum()),'detail':f'{len(g)} records; member={member}; archive/member date labels do not independently certify vintage; ID prefixes are not opening dates'})
                summary.append({'year':y,'category':cat,'check':'historical_vintage','status':'UNVERIFIED','count':len(g),'detail':'Preserved source archive/member labelled month; no authenticated publication/download timestamp establishes all row vintage'})
                if y==2020:
                    prefixes=g.facility_id.astype(str).str.slice(0,18).value_counts().head(5).to_dict()
                    summary.append({'year':y,'category':cat,'check':'historical_vintage_contrary_flags','status':'REVIEW','count':int(g.facility_id.astype(str).str.contains('2022').sum()),'detail':f'ID literal 2022 count (NOT an inferred opening year); frequent prefixes={prefixes}; nested archive path={member}. Source-row PASS does not establish source-vintage PASS.'})
                del raw,mapped,members
            elif cat=='CLINIC_PRIMARY':
                raw=clinic.reindex(g.facility_id.str.removeprefix('CLINIC_'))
                valid=raw['사업장명'].str.strip().to_numpy()==g.facility_name.to_numpy();valid&=raw['업태구분명'].str.strip().to_numpy()==g.facility_sub.to_numpy()
                target=f'{y-1}1231';date=lambda n:raw[n].str.replace('-','').str.strip()
                opening,closing,cancel,ss,se=[date(n) for n in ['인허가일자','폐업일자','인허가취소일자','휴업시작일자','휴업종료일자']]
                dateok=(opening.ne('')&(opening<=target)&(closing.eq('')|(closing>target))&(cancel.eq('')|(cancel>target))&~(ss.ne('')&se.ne('')&(ss<=target)&(se>=target))).to_numpy()
                x=pd.to_numeric(raw['좌표정보(X)'],errors='coerce');yy=pd.to_numeric(raw['좌표정보(Y)'],errors='coerce');lon,lat=trans.transform(x,yy)
                coords=np.isclose(lon,g.longitude,atol=1e-8,rtol=0)&np.isclose(lat,g.latitude,atol=1e-8,rtol=0)
                row.loc[mask,'source_row_match']=valid&dateok;row.loc[mask,'source_coordinate_match']=coords
                row.loc[mask,'coordinate_evidence_status']=np.where(coords,'MATCH_PRESERVED_TM_RAW','PENDING_GEOCODE_CHECK')
                non_tm=g.loc[~coords].set_index(['year_snapshot','facility_id']);geo=geocode.reindex(non_tm.index)
                gs=(non_tm[['facility_name','road_address','facility_sub']].fillna('').astype(str)==geo[['facility_name','road_address','facility_sub']].fillna('').astype(str)).all(axis=1)
                gc=np.isclose(non_tm.longitude,pd.to_numeric(geo.longitude),atol=1e-8,rtol=0)&np.isclose(non_tm.latitude,pd.to_numeric(geo.latitude),atol=1e-8,rtol=0)
                good=gs&gc&geo.geocode_status.ne('FAILED')
                good_ids=set(non_tm.index.get_level_values('facility_id')[good])
                row.loc[row.facility_id.isin(good_ids),'coordinate_evidence_status']='MATCH_PRESERVED_GEOCODE_LEDGER'
                summary.append({'year':y,'category':cat,'check':'non_tm_geocode_ledger_values','status':'PASS' if good.all() else 'REVIEW','count':int(good.sum()),'detail':f'{len(non_tm)} non-TM records: source year/id/name/address/subtype/final coordinates match saved successful ledger; API responses and historical positions not revalidated'})
                unresolved=int((ss.ne('')&se.eq('')&(ss<=target)).sum())
                summary.append({'year':y,'category':cat,'check':'full_raw_id_type_date_link','status':'PASS' if (valid&dateok).all() else 'REVIEW','count':int((valid&dateok).sum()),'detail':f'{len(g)} rows; raw TM coordinates match={sum(coords)}; suspension start without end={unresolved}. Current permit extract lacks full historical address/status event history.'})
            else:
                rawp,raw=school_roster(y) if cat=='SCHOOL_BASIC' else library_roster(y);inputs.add(rawp)
                keys=set(zip(raw.facility_name,raw.road_address));valid=np.array([(n,a) in keys for n,a in zip(g.facility_name,g.road_address)])
                row.loc[mask,'source_row_match']=valid
                if cat=='SCHOOL_BASIC':
                    est=raw.groupby(['facility_name','road_address']).establishment.first()
                    vals=[est.get((n,a),'') for n,a in zip(g.facility_name,g.road_address)];row.loc[mask,'establishment']=vals
                    typ=raw.groupby(['facility_name','road_address']).raw_school_type.first().map({'초등학교':'ELEMENTARY','중학교':'SECONDARY','고등학교':'SECONDARY'})
                    type_match=np.array([typ.get((n,a),'')==s for n,a,s in zip(g.facility_name,g.road_address,g.facility_sub)])
                    valid&=type_match;row.loc[mask,'source_row_match']=valid
                    detail=f'Oct 1 {y-1} roster; establishment={dict(Counter(vals))}; source school type mapping verified; original includes private schools'
                else:
                    valid&=g.facility_sub.eq('공공도서관').to_numpy();row.loc[mask,'source_row_match']=valid
                    detail=f'{y-1} annual roster public-general/children categories; name-based shared coordinate catalogue, historical address position unverified'
                summary.append({'year':y,'category':cat,'check':'full_roster_name_address_link','status':'PASS' if valid.all() else 'REVIEW','count':int(valid.sum()),'detail':f'{len(g)} rows; '+detail})
            summary.append({'year':y,'category':cat,'check':'historical_coordinate','status':'UNVERIFIED','count':len(g),'detail':'No complete dated address/entrance history validated. Current coordinates or present geocoding may differ from historical sites.'})
        row['analysis_included']=~row.facility_category.eq('SCHOOL_BASIC')|row.establishment.isin(['국립','공립'])
        all_rows.append(row)
    result=pd.DataFrame(summary);result.to_csv(C.RESULTS/'facility_provenance.csv',index=False,encoding='utf-8-sig')
    pd.concat(all_rows,ignore_index=True).to_parquet(C.RESULTS/'facility_lineage.parquet',index=False)
    pd.DataFrame(counts).to_csv(C.RESULTS/'facility_counts.csv',index=False,encoding='utf-8-sig')
    for p in [C.FACILITIES/'geocode_input_address.xlsx',C.RAW_FACILITIES/'school/seoul_school_locations_standard_20260320.json']:
        if p.exists():inputs.add(p)
    return result,inputs
