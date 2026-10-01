# -*- coding: utf-8 -*-
"""S5 점검: 원고(MD·DOCX·PDF)의 구조, 금지·어려운 표현, 약어 첫 정의, 표·그림 서식(JTG 양식), 핵심 수치와 결과표의 일치를 확인한다.
결과는 audit/s5_qa.json. 2026-10-02 개정: 캡션을 '번호 줄 + 제목 줄'(표)과 '번호(굵게) + 제목'(그림)으로 바꾼 뒤의 형식에 맞춤.
"""
import json, re
from pathlib import Path
import fitz
import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

PKG = Path(__file__).resolve().parents[1]
M = PKG / "manuscript"
md = (M / "연구2_생활권_구획_점검_4.2.md").read_text(encoding="utf-8")
doc = Document(next(M.glob("*.docx")))
pdf = fitz.open(next(M.glob("*.pdf")))
pdf_text = "\n".join(p.get_text() for p in pdf)

L2 = pd.read_csv(PKG / "results/tables/L2_city_summary_2025.csv").set_index("partition")
T2 = pd.read_csv(PKG / "results/reused_20260929/T2_city_summary.csv"); T2 = T2[T2["year"] == 2025].iloc[0]
T1 = pd.read_csv(PKG / "results/tables/T1_input_overview.csv").iloc[0]
checks = {}

# 1) 구조 ────────────────────────────────────────────────────────────────────────
checks["docx_tables_7"] = len(doc.tables) == 7
checks["docx_figures_5"] = len(doc.inline_shapes) == 5
checks["pdf_pages"] = len(pdf)
labels = [f"표 4.2-{i}" for i in range(1, 7)] + ["표 A-1"] + [f"그림 4.2-{i}." for i in range(1, 6)]
checks["pdf_has_all_captions"] = {c: (c in pdf_text) for c in labels}
checks["pdf_no_replacement_char"] = "�" not in pdf_text

# 2) 금지 표현과 어려운 낱말 ─────────────────────────────────────────────────────────
bad = ["JTG", "게재본", "March", "3월", "PASS", "BLOCKED", "검증 필요", "작업 메모", "TODO", "추후", "패키지", "기록 파일", "분이 걸렸다"]
hard = ["정합성", "연결요소", "결속도", "후향", "재표집", "co-association", "신장나무", "대리 규칙", "최빈 분할", "상충", "도출", "산출하였", "기술적 분류", "순차이", "특유"]
checks["forbidden_terms_found"] = [b for b in bad if b in md]
checks["hard_words_found"] = {h: md.count(h) for h in hard if h in md}

# 3) 약어는 처음 쓸 때 '전체 용어(이하, 약어)'로 정의 (요약과 본문에서 각각) ─────────────────
summary, body = md.split("## 4.2.1", 1)
body = body.split("## 참고문헌")[0]
def first_use_ok(text, abbr):
    m = re.search(rf"(?<![A-Za-z]){re.escape(abbr)}(?![A-Za-z])", text)
    if not m:
        return None                     # 이 부분에서 쓰지 않음
    return text[max(0, m.start() - 5):m.start()] == "(이하, "
abbrs = ["OD", "IFR", "IoU", "ARI", "HW·WH", "Q"]
checks["abbreviation_first_use"] = {"요약": {a: first_use_ok(summary, a) for a in abbrs}, "본문": {a: first_use_ok(body, a) for a in abbrs}}
abbr_ok = all(v in (True, None) for part in checks["abbreviation_first_use"].values() for v in part.values())
# 불일치 D(박사논문 공통 용어, 2026-10-02): 요약·본문에서 처음 쓸 때 정의, D_flow 표기 없음, 구별 동 수 범위는 결과표 값
L1 = pd.read_csv(PKG / "results/tables/L1_algorithm_comparison_2025.csv")
docx_text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
checks["d_term"] = {"summary_defines": "불일치(D)" in summary, "body_defines": "불일치(D)는 두 경계가" in body,
                    "no_D_flow": "D_flow" not in md and "D_flow" not in docx_text and "Dflow" not in docx_text,
                    "dong_range_from_L1": f"{int(L1.n_dong.min())}~{int(L1.n_dong.max())}개 동으로 된 구별 연결망" in md}
abbr_ok = abbr_ok and all(checks["d_term"].values())

# 4) 표 서식(JTG): 가로선 3줄(위·머리행 아래·아래), 세로선·좌우·안쪽 가로선·음영 없음, 머리행 보통 굵기 ─────
def tbl_ok(t):
    b = t._tbl.tblPr.find(qn("w:tblBorders"))
    if b is None:
        return False
    vals = {tag: (b.find(qn(f"w:{tag}")).get(qn("w:val")), b.find(qn(f"w:{tag}")).get(qn("w:sz"))) for tag in ("top", "bottom", "left", "right", "insideH", "insideV")}
    edges = vals["top"] == ("single", "4") and vals["bottom"] == ("single", "5") and all(vals[t][0] == "nil" for t in ("left", "right", "insideH", "insideV"))
    def head_bottom(c):
        tcb = c._element.tcPr.find(qn("w:tcBorders"))
        return tcb is not None and tcb.find(qn("w:bottom")) is not None and tcb.find(qn("w:bottom")).get(qn("w:sz")) == "4"
    head = all(head_bottom(c) for c in t.rows[0].cells)
    no_shade = not t._tbl.xpath(".//w:shd")
    head_plain = all(not r.bold for c in t.rows[0].cells for p in c.paragraphs for r in p.runs)
    return bool(edges and head and no_shade and head_plain)
checks["tables_jtg_rules"] = [tbl_ok(t) for t in doc.tables]
paras = {p._p: p for p in doc.paragraphs}
cap_ok = []
for t in doc.tables:                    # 표 캡션: 표 바로 앞 두 문단이 '번호(굵게)'와 '제목(보통)'
    prev = t._tbl.getprevious(); prev2 = prev.getprevious() if prev is not None else None
    p1, p2 = paras.get(prev2), paras.get(prev)
    cap_ok.append(bool(p1 is not None and p2 is not None and p1.text.startswith("표 ") and all(r.bold for r in p1.runs) and not any(r.bold for r in p2.runs)))
checks["table_captions_above_label_then_title"] = cap_ok
fig_ok = []
for i, p in enumerate(doc.paragraphs):  # 그림 캡션: 그림 문단 다음 문단이 '그림 번호.'(굵게) + 제목(보통)
    if p._p.xpath(".//w:drawing"):
        c = doc.paragraphs[i + 1]
        fig_ok.append(bool(c.runs and c.runs[0].text.startswith("그림 ") and c.runs[0].bold and len(c.runs) > 1 and not c.runs[1].bold))
checks["figure_captions_below_bold_label"] = fig_ok

# 5) 그림 해상도(Elsevier: 선화 1000 dpi, 혼합 500 dpi 이상) ───────────────────────────────
checks["figure_dpi"] = {p.name: round(Image.open(p).info.get("dpi", (0, 0))[0]) for p in sorted((PKG / "results/figures").glob("*.png"))}
dpi_ok = all(v >= 500 for v in checks["figure_dpi"].values())

# 6) 핵심 수치(본문 문자열 ↔ 결과표) ──────────────────────────────────────────────────
need = {
    "IoU": f"{T2['iou_1to1']:.3f}", "IFR 공식": f"{T2['ifr_lz'] * 100:.2f}%", "IFR Leiden": f"{T2['ifr_ld'] * 100:.2f}%", "D_flow": f"{T2['d_flow'] * 100:.2f}%",
    "Louvain 서울 IFR": f"{L2.loc['louvain', 'ifr_seoul_num_over_den'] * 100:.2f}", "행 수": f"{int(T1['selected_rows']):,}",
    "가장 잦은 결과 비율 Leiden": f"{L2.loc['leiden', 'modal_share_median'] * 100:.1f}", "가장 잦은 결과 비율 Louvain": f"{L2.loc['louvain', 'modal_share_median'] * 100:.1f}",
    "같은 결과 구 수": f"{int(L2.loc['leiden', 'districts_same_partition_leiden_louvain'])}개 구에서",
    "Louvain 결과 바뀐 구 수": f"Louvain {int(L2.loc['louvain', 'districts_stability_ari_below_1'])}개",
}
checks["numbers_in_text"] = {k: (v in md) for k, v in need.items()}

# 7) 수식: 모든 수식이 같은 크기(10.5 pt)·같은 영문 글꼴, 번호 (1)부터 순서대로, 본문이 '식 (n)'으로 가리킴, 표기용 기호가 남지 않음 ─────
eqs = [p for p in doc.paragraphs if p.text.startswith("\t") and re.search(r"\t\((\d+)\)$", p.text)]
eq_nums = [int(re.search(r"\((\d+)\)$", p.text).group(1)) for p in eqs]
eq_sizes = sorted({r.font.size.pt for p in eqs for r in p.runs if r.text.strip()})
eq_fonts = sorted({r.font.name for p in eqs for r in p.runs if r.text.strip() and not re.fullmatch(r"\t?\(\d+\)", r.text)})
all_text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
checks["equations"] = {"count": len(eqs), "numbers": eq_nums, "font_sizes_pt": eq_sizes, "latin_fonts": eq_fonts,
                       "referenced_in_text": {n: (f"식 ({n})" in md) for n in eq_nums},
                       "markup_left_in_docx": [m for m in ("_{", "^{", "*") if m in all_text],
                       "hangul_inside_equations": [p.text.strip() for p in eqs if re.search(r"[가-힣]", p.text)]}
checks["hyphen_minus_numbers"] = sorted(set(re.findall(r"(?<![0-9A-Za-z가-힣)\]])-\d[\d.,]*", all_text + "\n" + md)))  # 음수는 − (U+2212)
eq_ok = (not checks["hyphen_minus_numbers"] and eq_nums == list(range(1, len(eqs) + 1)) and len(eqs) == 5 and eq_sizes == [10.5] and eq_fonts == ["Times New Roman"]
         and all(checks["equations"]["referenced_in_text"].values()) and not checks["equations"]["markup_left_in_docx"]
         and not checks["equations"]["hangul_inside_equations"])
checks["equations"]["pass"] = bool(eq_ok)

checks["pass"] = bool(eq_ok and checks["docx_tables_7"] and checks["docx_figures_5"] and checks["pdf_no_replacement_char"] and all(checks["pdf_has_all_captions"].values())
                      and not checks["forbidden_terms_found"] and not checks["hard_words_found"] and abbr_ok and all(checks["tables_jtg_rules"])
                      and all(cap_ok) and all(fig_ok) and dpi_ok and all(checks["numbers_in_text"].values()))
(PKG / "audit" / "s5_qa.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(checks, ensure_ascii=False, indent=1))
assert checks["pass"]
