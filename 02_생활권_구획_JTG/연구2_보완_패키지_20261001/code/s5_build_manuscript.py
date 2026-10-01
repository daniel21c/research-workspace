# -*- coding: utf-8 -*-
"""S5-2 원고 생성: 학위논문 4.2절(한글)을 결과표에서 생성한다. 본문의 모든 수치는 CSV/JSON에서 읽는다(직접 입력 금지).
산출: manuscript/연구2_생활권_구획_점검_4.2.md, .docx  (PDF는 code/export_pdf.ps1 로 Word에서 출력)

2026-10-02 개정: (1) 약어는 처음 쓸 때 '전체 용어(이하, 약어)'로 정의(요약과 본문에서 각각 한 번), (2) 대학생이 읽을 수 있는 쉬운 어휘,
(3) 표·그림 서식을 JTG 게재본 양식에 맞춤: 표 캡션은 표 위에 '번호(굵게)' 줄과 '제목' 줄, 표 선은 가로선만(위·머리행 아래·아래),
세로선·좌우 테두리·안쪽 가로선·음영 없음, 머리행 보통 굵기, 주석은 표 아래 작은 글씨. 그림 캡션은 그림 아래 '그림 번호.'(굵게)+제목.
"""
from __future__ import annotations
import json, re
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

PKG = Path(__file__).resolve().parents[1]
RES = PKG / "results"; TAB = RES / "tables"; OLD = RES / "reused_20260929"; FIG = RES / "figures"
OUTD = PKG / "manuscript"; OUTD.mkdir(exist_ok=True)

# ── 표 서식(JTG 게재본에서 측정: 위 선 0.5pt, 머리행 아래 0.5pt, 아래 선 약 0.6pt, 세로선 없음) ─────────
RULE_TOP, RULE_HEAD, RULE_BOTTOM = 4, 4, 5          # Word 테두리 굵기 단위(1/8pt): 4=0.5pt, 5=0.625pt
TABLE_FONT_PT, TABLE_FONT_SMALL_PT = 9, 8
CAPTION_FONT_PT, NOTE_FONT_PT = 9.5, 8.5

# ── 입력 ─────────────────────────────────────────────────────────────────────────────
T1 = pd.read_csv(TAB / "T1_input_overview.csv").iloc[0]
C1 = pd.read_csv(TAB / "C1_official_zone_composition.csv")
L1 = pd.read_csv(TAB / "L1_algorithm_comparison_2025.csv")
L2 = pd.read_csv(TAB / "L2_city_summary_2025.csv").set_index("partition")
T2 = pd.read_csv(OLD / "T2_city_summary.csv"); T2 = T2[T2["year"] == 2025].iloc[0]
T3 = pd.read_csv(OLD / "T3_district_2025.csv")
T6 = pd.read_csv(OLD / "T6_sensitivity.csv"); T6 = T6[T6["year"] == 2025]
T7 = pd.read_csv(OLD / "T7_connected_reference_2025.csv")
T8 = pd.read_csv(OLD / "T8_correlations_2025.csv").set_index("y")
A2 = pd.read_csv(OLD / "A2_fixed2025_on2020.csv").iloc[0]
S = json.loads((OLD / "summary.json").read_text(encoding="utf-8"))
dong_leiden = pd.read_csv(PKG / "inputs" / "dong_to_leiden_2025_mapping_424.csv", encoding="utf-8-sig")
dong_louvain = pd.read_csv(PKG / "output" / "louvain" / "2025" / "metrics" / "louvain_mapping_2025.csv", encoding="utf-8-sig")
louv_info = json.loads((PKG / "output" / "louvain" / "2025" / "run_info.json").read_text(encoding="utf-8"))

# 9/29 결과표의 진단 구분 이름 → 원고에서 쓰는 쉬운 이름 (규칙은 같다)
SIGNAL = {"정합": "일치", "이동포착 증가·경계 검토": "이동 포착 증가·경계 검토", "상충·맥락 검토": "지표 엇갈림·맥락 검토"}


# ── 서식 도우미 ───────────────────────────────────────────────────────────────────────────
def p1(x): return f"{x * 100:.1f}"
def p2(x): return f"{x * 100:.2f}"
def _signed(v, nd):
    """부호를 붙인 수. 반올림해 0이면 부호 없이 0으로 쓴다(계산 오차로 생긴 -0.000·+0.00 표기를 막는다)."""
    s = f"{v:.{nd}f}"
    return s.lstrip("-") if float(s) == 0 else f"{v:+.{nd}f}"
def pp2(x): return _signed(x * 100, 2)
def n3(x): return f"{x:.3f}"
def sg3(x): return _signed(x, 3)
def n4(x): return f"{x:.4f}"
def sg4(x): return _signed(x, 4)
def cnt(x): return f"{int(round(x)):,}"
def sizes(df, ku, col): return "·".join(str(v) for v in sorted(df[df["Ku"] == ku].groupby(col).size().tolist(), reverse=True))
def names(df): return "·​".join(df["ku_name"].tolist())      # 영폭 공백: 긴 구 이름 목록에서 줄바꿈 위치를 허용


blocks = []
def H(level, text): blocks.append(("h", level, text))
def P(text): blocks.append(("p", text))
def EQ(text, num): blocks.append(("eq", text, num))      # 수식(표기 규칙은 rich()), num = 식 번호
def TBL(label, title, head, rows, note=None, small=False, widths=None, align=None):
    blocks.append(("t", label, title, head, rows, note, small, widths, align))
def REF(text): blocks.append(("r", text))
def FG(name, label, title, width=6.3): blocks.append(("f", name, label, title, width))


# ── 값 정리 ──────────────────────────────────────────────────────────────────────────────
le, lv, of = L2.loc["leiden"], L2.loc["louvain"], L2.loc["official"]
n_same = int(L2.loc["leiden", "districts_same_partition_leiden_louvain"])
diff = L1[L1["leiden_louvain_ari"] < 1 - 1e-12].copy()
louv_unstable = L1[L1["louvain_stability_ari_min"] < 1 - 1e-12].sort_values("louvain_stability_ari_min")
louv_modal_ne = L1[~L1["louvain_modal_equals_consensus"]]
low_modal = L1[L1["louvain_modal_share"] < 0.05].sort_values("louvain_modal_share")
ifr_louv_higher = L1[L1["leiden_minus_louvain_ifr"] < -1e-12]
q_louv_higher = L1[L1["leiden_minus_louvain_q"] < -1e-12]
big_q = diff.sort_values("leiden_minus_louvain_q")
sig = T3["review_signal"].map(SIGNAL).value_counts()
identical = T3[T3["iou_1to1"] > 1 - 1e-12]
low_iou = T3.sort_values("iou_1to1").head(4)
n_matched = int(round(T2["matched_share"] * 424))
g_pos, g_neg = int((T3["g"] > 1e-12).sum()), int((T3["g"] < -1e-12).sum())
ref_above = int((T7["observed"] > T7["q95"]).sum())
dq_flip = T6[(T6["delta_q"].abs() > 1e-9) & (np.sign(T6["delta_q"].round(9)) != np.sign(T6["delta_q_no_loop"].round(9)))]
g_flip = int((np.sign(T6["g"].round(10)) != np.sign(T6["g_no_loop"].round(10))).sum() + (np.sign(T6["g"].round(10)) != np.sign(T6["g_within_ku_den"].round(10))).sum())
tg = T3.set_index("ku_name")
cz = C1.groupby("Ku").size()
zone_dong = C1["n_dong"]
rho = lambda y: T8.loc[y]
mz, mr = L1[L1["ku_name"] == "중구"].iloc[0], L1[L1["ku_name"] == "중랑구"].iloc[0]


def modal_n_at_selected(row):   # 고른 해상도에서 가장 자주 나온 결과의 권역 수 (Louvain 해상도 기록)
    s = pd.read_csv(next((PKG / "output" / "louvain" / "2025" / "resolution_scan_logs").glob(f"{int(row['ku'])}_*_2025.csv")))
    return int(s[np.isclose(s["resolution"], row["louvain_resolution"])].iloc[0]["modal_n_communities"])


mz_n, mr_n = modal_n_at_selected(mz), modal_n_at_selected(mr)
assert mz_n > mz["target"] and mr_n > mr["target"]     # 본문 서술("목표보다 많은 권역")의 근거


# ═════════════════════════════════════ 본문 ═════════════════════════════════════════════
H(1, "4.2 생활권 구획과 공식 생활권 점검")

H(2, "이 절의 요약")
P(f"이 절은 이동 자료로 만든 권역과 서울시 공식 생활권 116개를 비교한다. 서울 424개 행정동과 구마다 정해진 공식 생활권 수를 그대로 두고, 2025년 1월 이동 자료로 구마다 권역을 나누었다. "
  f"첫째, 같은 절차로 두 가지 권역 찾기 방법(Louvain, Leiden)을 적용하자 둘 다 25개 구 모두에서 목표 권역 수를 맞추었고, 116개 권역이 모두 한 덩어리로 이어졌다. "
  f"두 방법의 결과는 {n_same}개 구에서 같고 {len(diff)}개 구에서 달랐으며, 같은 계산을 다시 했을 때 결과가 바뀌는 구는 Leiden {int(le['districts_stability_ari_below_1'])}개, Louvain {int(lv['districts_stability_ari_below_1'])}개였다. "
  f"둘째, 공식 생활권과 Leiden 권역이 얼마나 겹치는지를 본 Intersection over Union(이하, IoU)은 {n3(T2['iou_1to1'])}이었고, {len(identical)}개 구는 두 권역이 완전히 같았다. "
  f"셋째, 출발한 통행 가운데 같은 권역 안에서 끝난 통행의 비율인 Internal Flow Ratio(이하, IFR)는 공식 생활권 {p2(T2['ifr_lz'])}%, Leiden 권역 {p2(T2['ifr_ld'])}%로 {p2(T2['g'])}%p 차이였다. "
  f"그러나 두 권역 체계가 '권역 안'과 '권역 밖'을 서로 다르게 판정한 통행은 {p2(T2['d_flow'])}%로, 이 차이보다 훨씬 많았다. "
  f"넷째, 여러 지표를 함께 보면 먼저 살펴볼 구를 고를 수 있지만, 이것이 경계를 바꾸거나 공식 생활권을 대신하라는 뜻은 아니다.")

H(2, "4.2.1 연구의 목적과 위치")
P("생활권 같은 중간 계획 단위가 필요하다면, 다음 질문은 그 경계를 어떻게 정하고 어떻게 점검하느냐이다. 행정동은 자료를 모으고 책임을 나누기 쉬운 기본 단위이지만, 사람들의 일상 이동은 한 동 안에서 끝나지 않는다. "
  "반대로 자치구 하나를 통째로 한 권역으로 보면 구 안의 서로 다른 생활 반경을 구분하기 어렵다. 이 절은 동과 구 사이에 정해 둔 공식 생활권이 실제 이동을 얼마나 담아내는지, 그리고 그 결과를 계획 점검에 어떻게 쓸 수 있는지를 다룬다.")
P("이동 자료에서 서로 자주 오가는 동들의 묶음(커뮤니티)을 찾는 방법(community detection algorithm)으로 권역을 만들고, 이를 비교 기준(benchmark)으로 삼아 공식 생활권을 점검하는 방법은 Park et al.(2026)이 제시하였다. "
  "그 방법은 이동 자료 거르기, 구마다 목표 권역 수를 정한 권역 나누기, 여러 지표를 함께 보는 평가, 두 가지 권역 찾기 방법의 비교로 이루어진다. "
  "평가 지표는 구 안의 통행이 권역 안에 얼마나 몰려 있는지를 보는 Modularity(이하, Q), 출발한 통행 가운데 같은 권역 안에서 끝난 통행의 비율인 Internal Flow Ratio(이하, IFR), "
  "권역이 한 덩어리로 이어져 있는지를 보는 공간 연속성(spatial contiguity), 두 권역 체계가 얼마나 겹치는지를 보는 Intersection over Union(이하, IoU)이다. 이 절은 그 구성을 따르되, 모든 수치는 이 논문의 공통 자료와 공통 계산 절차로 새로 계산하였다.")
P("질문은 세 가지다. 첫째, 같은 424개 동과 구별 공식 생활권 수를 유지할 때 공식 생활권과 이동 권역은 얼마나 겹치며, 어디에서 어긋나는가. "
  "둘째, 같은 조건에서 Louvain과 Leiden은 목표 권역 수 맞추기, Q, IFR, 공간 연속성, 결과가 다시 나오는 정도에서 어떻게 다른가. 셋째, 지표를 함께 보면 어떤 구를 먼저 살펴볼 수 있으며, 그 해석의 한계는 무엇인가.")
P("공식 생활권은 행정 계획과 협의를 위한 단위이고, 이동 자료로 만든 권역은 이를 비교해 보는 기준이다. IFR이 높다고 서비스가 좋거나 공평하다는 뜻은 아니므로, 공식 생활권의 값이 낮다는 이유만으로 계획이 잘못되었다고 보지 않는다. "
  "이 절은 생활권 계획의 여러 목적 가운데 '실제 이동과 얼마나 맞는가'라는 한 가지 측면만 따로 살핀다. 경계를 조정했을 때 통행과 시설 이용이 어떻게 달라지는지는 4.3에서, 두 시점을 비교한 경계 동 점검은 4.4에서 다룬다.")

H(2, "4.2.2 자료와 표본")
P("분석 자료는 서울 생활이동 자료(서울특별시, 2021)를 행정동 사이의 통행량 표로 모은 2025년 1월 Origin–Destination(이하, OD) 자료이다. 도착 시각이 9시부터 20시대(09:00~20:59)인 통행만 쓰고, 요일은 가리지 않았다. "
  "자료 제공 기관이 집(Home)과 직장·학교(Work) 사이의 이동으로 분류한 유형(이하, HW·WH)은 뺐고, 출발지와 도착지가 모두 서울인 통행만 남겼다. "
  "HW·WH를 뺀 것은 일상생활 이동을 가려내기 위한 근사적인 기준이며, 남은 통행의 목적이 쇼핑이나 여가로 확인된 것은 아니다.")
P(f"행정동은 생활이동 자료의 동 코드에 맞춘 {int(T1['n_dong'])}개를 쓰고, 동 경계는 2023년 7월 1일 기준을 적용하였다. 원래 행정동 경계 파일에는 426개 구역이 있으나, 이동 자료에는 나뉘기 전 코드만 있는 오류2동·항동과 상일1동·상일2동을 각각 하나로 합쳤고, 개포3동은 코드를 바로잡았다. "
  f"공식 생활권은 서울시 2030 생활권계획의 지역생활권 116개(서울특별시, 2018)이다. 공식 생활권의 경계는 동 경계와 딱 맞지 않으므로, 각 동을 가장 많이 겹치는 공식 생활권에 넣어 동 단위의 공식 생활권을 만들었다. 따라서 이 결과가 원래 계획선을 그대로 옮긴 것은 아니다. "
  f"자치구는 {int(T1['n_district'])}개이고, 구마다 공식 생활권은 {int(cz.min())}~{int(cz.max())}개(합 {int(T1['n_official_zones'])}개), 생활권 하나에 속한 동은 중앙값 {zone_dong.median():.0f}개(범위 {int(zone_dong.min())}~{int(zone_dong.max())}개)이다.")
TBL("표 4.2-1", "분석 자료의 규모(2025년 1월).", ["", "값"], [
    ["행정동 / 자치구 / 공식 생활권", f"{int(T1['n_dong'])} / {int(T1['n_district'])} / {int(T1['n_official_zones'])}"],
    ["통행이 있는 OD 쌍", cnt(T1["od_pairs"])],
    ["조건에 맞는 원자료 행", cnt(T1["selected_rows"])],
    ["비공개 값이 있는 행 (비중)", f"{cnt(T1['masked_rows'])} ({p1(T1['masked_share_rows'])}%)"],
    ["전체 통행량", f"{T1['total_flow']:,.1f}"],
    ["같은 동 안 통행량 (비중)", f"{T1['within_dong_flow']:,.1f} ({p1(T1['within_dong_share_flow'])}%)"],
    ["같은 구 안 통행량 (비중)", f"{T1['within_district_flow']:,.1f} ({p1(T1['within_district_share_flow'])}%)"],
    ["같은 구 안 기록 행 (비중)", f"{cnt(T1['within_district_rows'])} ({p1(T1['within_district_share_rows'])}%)"],
], "주: 통행량은 이동 인구의 합이다. 행은 성별·연령·시간대 등으로 나뉜 자료의 기록 단위로, 통행량이나 사람 수와 다르다. 3명 미만이라 공개되지 않은 값은 0으로 두었다. 자료: 서울 생활이동 자료의 행정동 OD 집계표.")
P("같은 구 안 통행의 비중은 기록 행으로 세느냐, 통행량으로 세느냐에 따라 다르다. 이 절의 비율은 모두 통행량을 기준으로 하고, 행 수는 자료 규모를 보여 줄 때만 쓴다. "
  "OD의 출발지와 도착지를 서울로 한정했으므로 IFR의 분모는 서울 424개 동으로 가는 통행 전체이며, 서울 밖으로 가는 통행까지 넣은 자족 정도는 아니다. 통행량에는 같은 사람의 반복 이동이 들어 있으므로 주민 수로 읽지 않는다.")

H(2, "4.2.3 권역 나누기 절차와 지표")
H(3, "권역 나누기 절차")
P("권역은 자치구마다 따로 나눈다. 구 경계를 지키는 것이 행정적으로 실행하기 쉽기 때문이며, 구를 넘나드는 생활권을 찾는 것과는 다른 문제이다. "
  "구 안의 동을 점으로, 두 동 사이 양방향 통행량의 합(*f*_{*ij*}+*f*_{*ji*})을 두 점을 잇는 선의 무게로 두는 연결망을 만든다. 같은 동 안에서 끝난 통행(*f*_{*ii*})도 자기 자신으로 돌아오는 선으로 한 번 넣는다. 목표 권역 수는 그 구의 공식 생활권 수이다.")
P("권역 찾기 방법은 실행할 때마다 결과가 조금씩 달라질 수 있으므로, 다음 다섯 단계로 권역을 정한다.")
P("1단계: 권역을 잘게 나누는 정도를 정하는 해상도 값(resolution parameter, γ; Reichardt and Bornholdt, 2006)을 하나 정하고 권역 찾기를 3,000번 실행한다. 실행마다 난수의 시작값(시드)을 정해, 같은 설정이면 같은 결과가 나오게 한다.")
P("2단계: 3,000번 가운데 두 동이 같은 권역에 묶인 비율이 0.5 이상이면 두 동을 잇고, 이렇게 이어진 동의 묶음을 그 해상도의 합의 권역(consensus partition)으로 삼는다(Lancichinetti and Fortunato, 2012).")
P("3단계: 해상도를 0.01부터 2.50까지 0.01씩 바꾸며 250가지에 대해 1~2단계를 되풀이한다.")
P("4단계: 권역 수가 목표와 같은 해상도 가운데 Q가 가장 큰 것을 고른다. Q가 같으면 IFR이 큰 것, 그래도 같으면 해상도가 낮은 것을 고른다.")
P("5단계: 고른 해상도에서 1~2단계를 10번 더 되풀이해 같은 권역이 다시 나오는지 확인한다. 두 결과가 얼마나 같은지는 Adjusted Rand Index(이하, ARI; Hubert and Arabie, 1985)로 잰다. ARI는 1이면 완전히 같고, 0에 가까우면 우연히 맞는 수준이다.")
P("권역 찾기 방법으로는 Leiden(Traag et al., 2019)과 Louvain(Blondel et al., 2008)을 쓴다. Leiden은 Louvain에서 생길 수 있는 '서로 이어지지 않은 동이 한 권역에 묶이는 문제'를 보완한 방법이다. "
  "두 방법은 위 절차에서 권역을 찾는 계산 한 단계만 다르고, 연결망, 해상도 범위, 반복 횟수, 합의 기준, 선택 규칙, 시드 규칙은 같다. Leiden 권역은 논문 공통 계산에서 확정한 결과를 그대로 쓰고, Louvain 권역은 같은 코드에서 그 한 단계만 바꾸어 새로 만들었다.")
FG("F1_procedure", "그림 4.2-1.", "이동 권역을 만들고 공식 생활권을 점검하는 절차. (a)는 구마다 이동 권역을 만드는 단계, (b)는 그 권역으로 공식 생활권을 점검하는 단계이다. OD: Origin–Destination, HW·WH: 집과 직장·학교 사이 이동 유형, Q: Modularity, IFR: Internal Flow Ratio, IoU: Intersection over Union, ARI: Adjusted Rand Index.")
TBL("표 4.2-2", "권역 나누기 조건(Leiden과 Louvain에 같게 적용).", ["", "조건"], [
    ["분석 단위", "자치구마다 따로 만든 연결망. 점은 동(모두 424개)"],
    ["선의 무게", "두 동 사이 양방향 통행량의 합(*f*_{*ij*}+*f*_{*ji*}). 같은 동 안 통행(*f*_{*ii*})은 자기 자신으로 돌아오는 선 1개"],
    ["해상도 값(γ)", "0.01~2.50, 0.01 간격(250가지)"],
    ["반복과 시드", "해상도마다 3,000번. 시드는 구·해상도·반복 번호로 정함"],
    ["합의 권역", "두 동이 같은 권역에 묶인 비율이 0.5 이상인 동끼리 이은 묶음"],
    ["목표 권역 수", "구별 공식 생활권 수(합 116)"],
    ["해상도 고르기", "목표 권역 수를 맞춘 해상도 중 Q가 가장 큰 것. 같으면 IFR, 그다음 낮은 해상도"],
    ["다시 나오는지 확인", "고른 해상도에서 10번 더 계산해 처음 결과와의 ARI를 봄"],
], "주: Leiden은 논문 공통 계산의 확정 결과를 쓰고, Louvain은 같은 코드에서 권역을 찾는 계산 한 단계만 바꾸어 만들었다.", widths=[3.6, 12.4])

H(3, "지표")
P("*f*_{*ij*}는 동 *i*에서 동 *j*로 간 통행량, *V*는 서울의 424개 동 전체, *i*∈*k*는 동 *i*가 구 *k*에 속한다는 뜻, *T*_{*k*}는 구 *k*의 동에서 출발해 서울 안 어디로든 간 통행량이다. 모든 IFR은 방향이 있는 OD에서 출발지를 기준으로 계산한다. 구나 서울 전체의 값은 비율을 평균하지 않고, 분자와 분모를 각각 더한 뒤 나눈다.")
EQ("IFR_{*k*}(*P*) = Σ_{*i*∈*k*, *j*∈*V*} *f*_{*ij*} · 1[*P*_{*i*} = *P*_{*j*}] / *T*_{*k*}", 1)
P("IFR은 식 (1)로 계산하며, 권역이 생활 이동을 안에 얼마나 담는지를 나타낸다. 기능 지역 연구에서 쓰는 자족성(self-containment)을 출발지 기준으로 잰 값에 해당한다. *P*_{*i*}는 동 *i*가 속한 권역이고, 1[ ]은 괄호 안이 참이면 1, 아니면 0이다.")
EQ("*G*_{*k*} = IFR_{*k*}(*L*) − IFR_{*k*}(*Z*)", 2)
EQ("*D*_{flow,*k*} = (*a*_{*k*} + *b*_{*k*}) / *T*_{*k*}", 3)
P("G는 식 (2)처럼 이동 권역(*L*)의 IFR에서 공식 생활권(*Z*)의 IFR을 뺀 값으로, 양수이면 이동 권역이 통행을 더 많이 안에 담는다는 뜻이다. D_flow는 식 (3)의 값이다. *a*_{*k*}는 이동 권역에서만 '권역 안'인 통행량, *b*_{*k*}는 공식 생활권에서만 '권역 안'인 통행량이며, D_flow는 한 체계에서만 '권역 안'으로 판정되는 통행의 비율로, 두 체계가 통행을 얼마나 다르게 나누는지를 보여 준다. "
  "분모가 같으므로 D_flow는 G의 크기(절댓값)보다 작을 수 없다. 서로 반대 방향의 차이가 맞비겨지면 G가 0이어도 D_flow는 클 수 있다. D_flow는 경계를 넘는 모든 통행의 비율이 아니다.")
EQ("*Q*_{*k*} = Σ_{*c*} [ *I*_{*c*} / *F*_{*k*} − ((*O*_{*c*} + *D*_{*c*}) / 2*F*_{*k*})^{2} ]", 4)
P("Q는 식 (4)로 계산하며, 구 안의 통행이 권역 안에 얼마나 몰려 있는지를 통행이 무작위로 흩어졌을 때 기대되는 값과 비교한 것이다(Newman, 2006; 해상도 1 기준). 값이 클수록 권역 안 이동이 상대적으로 많다. "
  "*I*_{*c*}는 권역 *c* 안의 통행량, *O*_{*c*}와 *D*_{*c*}는 권역 *c*에서 나가고 들어오는 통행량, *F*_{*k*}는 구 안 통행량의 합이다. 구마다 연결망이 다르므로 구별 Q를 평균해 '서울의 Q'라고 하지 않고, 중앙값과 범위로 보고한다.")
EQ("*S*_{*k*} = max_{π} Σ_{*z*} |*z* ∩ π(*z*)|,     IoU_{*k*} = *S*_{*k*} / (2*N*_{*k*} − *S*_{*k*})", 5)
P("IoU는 두 집합이 겹치는 정도를 재는 Jaccard 지수(Jaccard, 1912)를 권역 짝에 적용한 것으로, 식 (5)로 계산한다. 두 체계의 권역을 1:1로 짝짓되(π), 짝지은 권역끼리 겹치는 동 수의 합(*S*)이 가장 커지도록 짝을 정한다. *N*_{*k*}는 구의 동 수이다. 서울 전체 값도 구별 IoU의 평균이 아니라 *S*와 *N*을 모두 더해 계산한다. "
  "짝을 짓지 않고 두 결과를 비교하는 ARI도 함께 본다. 공식 생활권마다 가장 많이 겹치는 이동 권역을 각각 골라 겹친 동의 비율을 구한 값은 여러 생활권이 같은 권역을 고를 수 있으므로, IoU와 섞지 않고 '다대일 일치율'로 따로 적는다. "
  "공간 연속성은 권역에 속한 동들이 서로 맞닿아 한 덩어리를 이루는지를 본 것이며, 걸어서나 교통수단으로 이어진다는 뜻은 아니다.")

H(2, "4.2.4 Louvain과 Leiden의 비교")
P(f"같은 조건에서 두 방법은 서울 전체로 {int(le['n_communities_total'])}개 권역을 만들었고, 25개 구 모두에서 목표 권역 수를 맞추었으며, {int(le['contiguous_communities'])}개 권역이 모두 한 덩어리로 이어졌다(표 4.2-3). "
  f"목표 권역 수와 연속성만 보면 두 방법을 구별할 수 없다. 차이는 권역의 내용과, 같은 계산을 다시 했을 때 결과가 얼마나 그대로 나오는지에서 나타났다.")
TBL("표 4.2-3", "공식 생활권, Leiden 권역, Louvain 권역의 성질(2025년).", ["", "공식 생활권", "Leiden", "Louvain"], [
    ["서울 전체 권역 수", cnt(of["n_communities_total"]), cnt(le["n_communities_total"]), cnt(lv["n_communities_total"])],
    ["목표 권역 수를 맞춘 구", "–", f"{int(le['districts_hit_target'])} / 25", f"{int(lv['districts_hit_target'])} / 25"],
    ["한 덩어리로 이어진 권역", f"{int(of['contiguous_communities'])} / {int(of['n_communities_total'])}", f"{int(le['contiguous_communities'])} / {int(le['n_communities_total'])}", f"{int(lv['contiguous_communities'])} / {int(lv['n_communities_total'])}"],
    ["구별 Q, 중앙값 (범위)", f"{n4(of['q_district_median'])} ({n4(of['q_district_min'])}–{n4(of['q_district_max'])})", f"{n4(le['q_district_median'])} ({n4(le['q_district_min'])}–{n4(le['q_district_max'])})", f"{n4(lv['q_district_median'])} ({n4(lv['q_district_min'])}–{n4(lv['q_district_max'])})"],
    ["서울 전체 IFR (%)", p2(of["ifr_seoul_num_over_den"]), p2(le["ifr_seoul_num_over_den"]), p2(lv["ifr_seoul_num_over_den"])],
    ["가장 자주 나온 결과의 비율, 중앙값 (%)", "–", p1(le["modal_share_median"]), p1(lv["modal_share_median"])],
    ["합의 권역이 가장 자주 나온 결과와 다른 구", "–", f"{int(le['districts_modal_not_equal_consensus'])}", f"{int(lv['districts_modal_not_equal_consensus'])}"],
    ["다시 계산해 결과가 바뀐 구 (ARI 최솟값)", "–", f"{int(le['districts_stability_ari_below_1'])} ({n3(le['stability_ari_min_overall'])})", f"{int(lv['districts_stability_ari_below_1'])} ({n3(lv['stability_ari_min_overall'])})"],
], "주: 구별 Q는 구 안 연결망에서 해상도 1로 계산하였다. 서울 전체 IFR은 구별 분자의 합을 분모의 합으로 나눈 값이다. '가장 자주 나온 결과의 비율'은 고른 해상도의 3,000번 실행 가운데 같은 결과가 나온 비율이다.", widths=[6.6, 3.1, 3.1, 3.1])
P(f"두 방법의 결과는 {n_same}개 구에서 완전히 같았고, {len(diff)}개 구({names(diff)})에서 달랐다. 서울 전체 IFR은 Leiden {p2(le['ifr_seoul_num_over_den'])}%, Louvain {p2(lv['ifr_seoul_num_over_den'])}%로 거의 같지만 Q는 다르다. "
  f"구별 Q가 Leiden에서 Louvain과 같거나 큰 구는 {int(L2.loc['leiden', 'districts_leiden_q_ge_louvain'])}개이고, Louvain의 Q가 더 큰 구({names(q_louv_higher)})에서도 그 차이는 {', '.join(f'{abs(v):.4f}' for v in q_louv_higher['leiden_minus_louvain_q'])}에 그쳤다. "
  f"반대로 Leiden의 Q가 크게 높은 구는 {big_q.iloc[-1]['ku_name']}(Leiden {n4(big_q.iloc[-1]['leiden_q'])}, Louvain {n4(big_q.iloc[-1]['louvain_q'])})와 {big_q.iloc[-2]['ku_name']}(Leiden {n4(big_q.iloc[-2]['leiden_q'])}, Louvain {n4(big_q.iloc[-2]['louvain_q'])})이다. "
  f"서울 전체 IFR이 Louvain에서 조금 높은 것은 {names(ifr_louv_higher)}에서 Louvain의 IFR이 더 높기 때문이며({', '.join(f'{abs(v) * 100:.2f}%p' for v in ifr_louv_higher['leiden_minus_louvain_ifr'])}), 이 가운데 중구와 중랑구는 Q가 크게 낮다. "
  "IFR이 높은 권역이 Q는 낮을 수 있으므로, 한 가지 지표만으로 권역을 평가하지 않는다.")
FG("F2_maps_2025", "그림 4.2-2.", "공식 생활권(a), Leiden 권역(b), Louvain 권역(c), 2025년. 굵은 선은 자치구 경계, 가는 선은 권역 경계이며, 색은 같은 구 안의 권역을 구분하기 위한 것으로 값의 크기를 뜻하지 않는다. 지도의 선은 연구 지역의 구분을 나타내며 공인된 경계를 뜻하지 않는다. 자료: 2023년 7월 기준 행정동 경계, 서울시 2030 생활권계획, 서울 생활이동 자료(2025년 1월).")
rows = []
for r in diff.sort_values("leiden_minus_louvain_q", ascending=False).itertuples():
    rows.append([r.ku_name, n4(r.official_q), n4(r.leiden_q), n4(r.louvain_q), p2(r.leiden_ifr), p2(r.louvain_ifr), sizes(dong_leiden, r.ku, "community"), sizes(dong_louvain, r.ku, "community"), n3(r.leiden_louvain_ari)])
TBL("표 4.2-4", "두 방법의 결과가 다른 구.", ["구", "Q 공식", "Q Leiden", "Q Louvain", "IFR Leiden (%)", "IFR Louvain (%)", "권역별 동 수 Leiden", "권역별 동 수 Louvain", "두 결과의 ARI"], rows,
    "주: 권역별 동 수는 구 안의 각 권역에 속한 동의 수를 큰 순서로 적은 것이다. 두 결과가 같은 18개 구는 뺐고, 25개 구 전체는 부록 표 A-1에 있다.", small=True,
    widths=[1.5, 1.4, 1.5, 1.6, 1.7, 1.8, 2.5, 2.5, 1.5])
P(f"결과가 다시 나오는 정도에서는 차이가 더 뚜렷하다(그림 4.2-3). 고른 해상도의 3,000번 실행 가운데 가장 자주 나온 결과의 비율은 중앙값이 Leiden {p1(le['modal_share_median'])}%, Louvain {p1(lv['modal_share_median'])}%였다. "
  f"Louvain에서는 이 비율이 5%도 안 되는 구가 {len(low_modal)}개({names(low_modal)})였다. 합의 권역이 가장 자주 나온 결과와 다른 구도 Louvain에서만 {int(lv['districts_modal_not_equal_consensus'])}개({names(louv_modal_ne)}) 나타났다. "
  f"고른 해상도에서 같은 계산을 10번 다시 했을 때 처음과 다른 권역이 나온 구(ARI 1 미만)는 Leiden {int(le['districts_stability_ari_below_1'])}개, Louvain {int(lv['districts_stability_ari_below_1'])}개"
  f"({', '.join(f'{r.ku_name} {n3(r.louvain_stability_ari_min)}' for r in louv_unstable.itertuples())})였다.")
FG("F3_algorithm_comparison", "그림 4.2-3.", "구별 Q(a)와 결과가 다시 나오는 정도(b), 2025년. (a)의 선은 같은 구의 Leiden과 Louvain 값을 잇고, 빈 원은 공식 생활권 값이다. 구는 두 방법의 Q 차이가 큰 순서로 위에서부터 놓았다. (b)는 고른 해상도의 3,000번 실행 가운데 가장 자주 나온 결과의 비율이며, Louvain에서 같은 계산을 10번 다시 해 결과가 바뀐 구에는 ARI(Adjusted Rand Index) 최솟값을 적었다. Q: Modularity. 자료: 서울 생활이동 자료(2025년 1월)로 계산.")
P(f"중구와 중랑구는 이 차이가 가장 크게 나타난 구이다. Louvain은 중구에서 γ={mz['louvain_resolution']:.2f}, 중랑구에서 γ={mr['louvain_resolution']:.2f}라는 낮은 해상도에서만 목표 권역 수 {int(mz['target'])}개를 맞추었고, 목표 수를 맞춘 해상도는 각각 {int(mz['louvain_n_resolutions_hitting_target'])}개, {int(mr['louvain_n_resolutions_hitting_target'])}개뿐이었다. "
  f"이 해상도에서 한 번 한 번의 실행은 목표보다 많은 {mz_n}개와 {mr_n}개 권역을 가장 자주 냈고, 그 결과가 나온 비율도 {p1(mz['louvain_modal_share'])}%와 {p1(mr['louvain_modal_share'])}%에 그쳤다. 목표 권역 수는 3,000번의 결과를 모아 합의 권역을 만든 뒤에야 맞추어졌다. "
  f"이렇게 만든 Louvain 권역의 Q는 각각 {n4(mz['louvain_q'])}, {n4(mr['louvain_q'])}로 Leiden의 {n4(mz['leiden_q'])}, {n4(mr['leiden_q'])}보다 낮다. 중랑구에서 Louvain의 가장 큰 권역은 동 {sizes(dong_louvain, 11070, 'community').split('·')[0]}개로 Leiden의 {sizes(dong_leiden, 11070, 'community').split('·')[0]}개보다 크다. "
  f"Leiden은 두 구에서 γ={mz['leiden_resolution']:.2f}, {mr['leiden_resolution']:.2f}로 목표 권역 수를 맞추었다. 이런 차이가 Louvain의 어떤 성질에서 생기는지는 이 분석에서 밝히지 않았다.")
P("이 비교는 15~27개 동으로 된 구별 연결망, 구별 목표 권역 수, 3,000번 결과를 모은 합의 권역이라는 조건에서 나온 것이므로, 두 방법 가운데 어느 쪽이 언제나 낫다는 뜻은 아니다. "
  "다만 같은 절차를 적용해도 Leiden의 권역이 다시 계산할 때 더 그대로 나오고, 목표 권역 수만으로는 드러나지 않는 권역의 질 차이가 일부 구에서 나타났다. 이것이 이후 공식 생활권과의 비교에 Leiden 권역을 쓰는 이유이다.")

H(2, "4.2.5 공식 생활권과 이동 권역의 겹침")
P(f"공식 생활권과 Leiden 권역의 구별 IoU는 중앙값 {n3(S['median_iou'])}, 범위 {n3(S['range_iou'][0])}~{n3(S['range_iou'][1])}이다. 서울 전체에서 두 체계가 같은 권역으로 짝지은 동은 {n_matched}개로 424개 동의 {p2(T2['matched_share'])}%이며, "
  f"이를 IoU로 나타내면 {n3(T2['iou_1to1'])}, ARI는 {n3(T2['ari'])}, 다대일 일치율은 {p2(T2['jaccard_maxmatch_share'])}%이다. 계산 방법과 분모가 다른 값들이므로 모두 '일치율'로 부르지 않고 이름을 나누어 쓴다.")
P(f"두 체계가 완전히 같은 구는 {len(identical)}개({names(identical)})이다. 반면 {', '.join(f'{r.ku_name}({n3(r.iou_1to1)})' for r in low_iou.itertuples())}는 겹침이 적다. 공간 연속성은 두 체계 모두 116개 권역이 이어져 있어 차이가 없다.")
P(f"구와 권역 수를 같게 맞추면 우연만으로도 어느 정도는 겹친다. 이 정도를 가늠하려고, 구마다 서로 맞닿은 동끼리 무작위로 묶어 목표 권역 수만큼 나눈 '비교용 무작위 권역'을 {int(T7['draws'].iloc[0]):,}개씩 만들었다. "
  f"무작위로 동을 하나씩 이어 구 전체를 잇는 나무 모양의 연결을 만든 뒤, 연결 몇 개를 끊어 권역으로 나누는 방식이다. 관측된 IoU가 무작위 권역의 상위 5% 경계(95번째 백분위수)보다 큰 구는 {ref_above}개였다. "
  "다만 이 방식은 가능한 모든 무작위 권역을 고르게 뽑은 것이 아니고 권역의 크기도 맞추지 않았으므로, 이 결과를 일반적인 '우연일 확률'이나 계획 경계의 통계적 근거로 읽지 않는다.")

H(2, "4.2.6 IFR의 차이와 판정이 바뀌는 통행")
P(f"서울 전체에서 공식 생활권의 IFR은 {p2(T2['ifr_lz'])}%, Leiden 권역은 {p2(T2['ifr_ld'])}%로 G는 {pp2(T2['g'])}%p이다. 그러나 공식 생활권에서만 '권역 안'인 통행이 전체의 {p2(T2['lz_only'])}%, Leiden 권역에서만 '권역 안'인 통행이 {p2(T2['ld_only'])}%여서, "
  f"두 체계가 안과 밖을 다르게 판정한 통행(D_flow)은 {p2(T2['d_flow'])}%이다. G는 이 두 몫이 서로 맞비기고 남은 차이이다. 구별로 보면 G가 양수인 구가 {g_pos}개, 음수인 구가 {g_neg}개, 0인 구가 {int(S['zero_g'])}개이다.")
rows = []
for r in T3.itertuples():
    rows.append([r.ku_name, f"{int(r.n_dong)}/{int(r.n_zones)}", n3(r.iou_1to1), p2(r.ifr_lz), p2(r.ifr_ld), pp2(r.g), p2(r.d_flow), sg4(r.delta_q), SIGNAL[r.review_signal]])
TBL("표 4.2-5", "구별 공식 생활권과 Leiden 권역의 비교(2025년).", ["구", "동/권역", "IoU", "IFR 공식 (%)", "IFR Leiden (%)", "G (%p)", "D_flow (%)", "ΔQ", "구분"], rows,
    "주: ΔQ는 Leiden 권역의 Q에서 공식 생활권의 Q를 뺀 값이다. 구분은 IoU가 1이면 '일치', G와 ΔQ가 모두 양수이면 '이동 포착 증가·경계 검토', 그 밖은 '지표 엇갈림·맥락 검토'이다. 이 구분은 설명을 위한 것이며 순위나 처방이 아니다.", small=True,
    widths=[1.6, 1.5, 1.3, 1.7, 1.8, 1.4, 1.6, 1.4, 3.7])
FG("F4_district_diagnostics", "그림 4.2-4.", "구별 IoU(a), IFR 차이 G(b), 판정이 바뀌는 통행의 비율 D_flow(c), 2025년. 구는 IoU가 작은 순서로 아래에서부터 놓았다. G = IFR(Leiden 권역) − IFR(공식 생활권)이며 G만 양과 음의 방향이 있다. IoU: Intersection over Union, IFR: Internal Flow Ratio. 자료: 서울 생활이동 자료(2025년 1월)로 계산.")
P("IFR은 이 자료에서 출발지를 기준으로 본 '권역 안에서 끝난 통행의 비율'이다. 같은 사람의 반복 이동도 따로 세므로 주민의 만족도나 생활 반경의 다양성과는 다르다. 같은 동 안의 통행은 두 체계에서 모두 '권역 안'이므로 두 IFR을 함께 높이지만, 판정이 바뀌는 통행에는 들어가지 않는다.")

H(2, "4.2.7 지표를 함께 읽기와 후속 조사 후보")
P(f"구별 지표를 함께 보면 서로 다른 상황이 구별된다. 25개 구는 '일치' {int(sig.get('일치', 0))}개, '이동 포착 증가·경계 검토' {int(sig.get('이동 포착 증가·경계 검토', 0))}개, '지표 엇갈림·맥락 검토' {int(sig.get('지표 엇갈림·맥락 검토', 0))}개로 나뉜다. "
  f"'지표 엇갈림·맥락 검토'에 속한 구는 {names(T3[T3['review_signal'].map(SIGNAL) == '지표 엇갈림·맥락 검토'])}이다. 이 구분은 순위를 매기거나 처방을 내리는 기준이 아니라, 어디부터 자세히 볼지를 정하는 출발점이다.")
gj, gd = tg.loc["광진구"], tg.loc["강동구"]; yc = tg.loc["양천구"]; jr, yd = tg.loc["종로구"], tg.loc["영등포구"]; sp = tg.loc["송파구"]
P(f"광진구와 강동구는 IoU가 가장 낮고({n3(gj['iou_1to1'])}, {n3(gd['iou_1to1'])}) D_flow가 가장 크다({p2(gj['d_flow'])}%, {p2(gd['d_flow'])}%). G는 {pp2(gj['g'])}%p, {pp2(gd['g'])}%p로 양수여서, 두 체계가 많은 통행을 다르게 묶고 Leiden 권역이 통행을 더 많이 안에 담는다. 후속 조사에서 먼저 살펴볼 구이다.")
P(f"양천구는 G가 {pp2(yc['g'])}%p로 거의 0이지만 D_flow는 {p2(yc['d_flow'])}%이다. 반대 방향의 차이가 거의 맞비겨져서 G만 보면 두 체계가 같은 이동을 담는 것처럼 보이지만, 실제로는 많은 통행의 판정이 바뀐다. 따라서 G가 작다는 것만으로 두 경계가 같다고 볼 수 없다.")
P(f"종로구와 영등포구는 G가 {pp2(jr['g'])}%p, {pp2(yd['g'])}%p로 음수이지만 ΔQ는 {sg4(jr['delta_q'])}, {sg4(yd['delta_q'])}로 양수이다. 구 안 통행이 권역 안에 몰린 정도(Q)와 서울 전체 목적지를 분모로 한 IFR이 서로 다른 방향을 가리킨 것으로, 두 지표가 서로 다른 질문에 답한다는 것을 보여 준다. "
  f"송파구는 G가 {pp2(sp['g'])}%p로 가장 크고 D_flow가 {p2(sp['d_flow'])}%이다. 이런 결과가 특정 시설이나 아파트 단지, 교통망 때문인지는 이 분석으로 알 수 없다.")
P("후속 검토에서 경계를 바꾸는 것만이 대응은 아니다. 경계가 나누고 있는 생활 반경을 확인한 뒤, 이웃한 생활권이 서비스를 함께 운영하거나 서로의 연결을 개선하는 것도 방법이 될 수 있다. 이 절은 그런 대응의 효과를 추정하지 않았다.")

H(2, "4.2.8 지표 사이의 관계와 평가 방식을 바꾼 결과")
FG("F5_iou_relations", "그림 4.2-5.", "구별 IoU와 G(a), ΔQ(b), D_flow(c)의 관계, 2025년. 점 하나가 구 하나이다. ρ는 스피어만 순위상관계수이고, 괄호는 25개 구를 중복을 허용해 다시 뽑아 5,000번 계산한 95% 범위이다. ΔQ = Q(Leiden 권역) − Q(공식 생활권). 자료: 서울 생활이동 자료(2025년 1월)로 계산.")
TBL("표 4.2-6", "공간적 겹침(IoU)과 이동 지표의 구별 관계.", ["", "ρ", "95% 범위"], [
    ["IoU와 G", f"{rho('g')['rho']:.3f}", f"[{rho('g')['lo']:.3f}, {rho('g')['hi']:.3f}]"],
    ["IoU와 ΔQ", f"{rho('delta_q')['rho']:.3f}", f"[{rho('delta_q')['lo']:.3f}, {rho('delta_q')['hi']:.3f}]"],
    ["IoU와 D_flow", f"{rho('d_flow')['rho']:.3f}", f"[{rho('d_flow')['lo']:.3f}, {rho('d_flow')['hi']:.3f}]"],
], "주: ρ는 스피어만 순위상관계수이다. 95% 범위는 25개 구를 중복을 허용해 다시 뽑아(부트스트랩) 5,000번 계산한 결과이며, 구 사이의 공간적 의존을 고려한 추론이나 인과 효과가 아니다.", widths=[5.0, 3.0, 5.0])
P(f"IoU와 G의 순위상관은 {rho('g')['rho']:.3f}이고 95% 범위가 0을 포함한다. 따라서 '겹침이 적을수록 G가 크다'고 일반화하기는 어렵다. IoU와 ΔQ, D_flow는 강한 음의 관계를 보이지만, D_flow는 권역이 달라진 동 사이의 통행으로 만들어지므로 두 지표는 계산 구조상 이어져 있다. "
  "또 이동 권역은 같은 이동 자료에서 Q가 커지도록 골랐으므로 ΔQ 비교는 이동 권역에 유리하다. 이 결과는 이동 자료에 얼마나 잘 맞는지를 비교한 것이지, 계획 전체가 낫다거나 서비스 성과가 좋다는 비교가 아니다.")
P(f"이미 정한 권역은 그대로 두고 평가 방식만 바꾸어 보았다. 같은 동 안 통행을 빼거나 분모를 같은 구 안 도착 통행으로 좁혀도 G의 부호가 바뀐 구는 {g_flip}개였다. 다만 두 체계에 공통인 분자나 분모를 함께 바꾼 계산이므로, 결과가 튼튼하다는 독립적인 증거로 쓰지는 않는다. "
  f"Q는 다르다. 같은 동 안 통행을 빼면 {'·'.join(dq_flip['ku_name'].tolist())}에서 ΔQ의 부호가 바뀌므로, 'Leiden 권역이 Q에서 언제나 앞선다'는 해석은 같은 동 안 통행을 넣는다는 정의를 전제로 한다.")
P(f"2025년 권역을 그대로 두고 2020년 1월 OD에 적용하면 공식 생활권 IFR {p2(A2['ifr_lz'])}%, Leiden 권역 IFR {p2(A2['ifr_ld'])}%, G {pp2(A2['g'])}%p, D_flow {p2(A2['d_flow'])}%이다. "
  "2020년 이동에서도 같은 방향의 차이가 보이지만, 같은 자료 체계와 같은 동 경계를 쓴 과거 자료에 적용한 것이므로 미래 예측이나 외부 검증은 아니다.")

H(2, "4.2.9 논의, 한계와 소결")
P("첫째, 공식 생활권은 이동과 상관없이 그은 임의의 선이 아니다. 여러 구에서 이동 권역과 많이 겹치고 일부 구는 완전히 같다. 다만 이 겹침의 일부는 같은 동과 구 경계, 같은 권역 수를 쓴 데서 생기므로 비교용 무작위 권역과 함께 읽어야 한다. "
  "또 공식 계획이 고려한 맥락이 모두 이동 자료에 담겼다고 할 수는 없다.")
P("둘째, 차이의 크기와 방향을 나누어 보면 서로 다른 상황을 구별할 수 있다. IoU가 낮고 G가 큰 곳은 권역 안에 담기는 이동이 늘어나는지 살펴볼 곳이고, D_flow는 큰데 G가 작은 곳은 반대 방향의 차이가 서로 맞비기는 곳이며, Q와 G가 다른 방향인 곳은 각 지표가 무엇을 재는지 확인한 뒤에야 판단할 수 있다. "
  "진단의 단위인 구는 정책을 결정하는 단위와 같지 않으며, 구별 비교는 같은 분모와 보고 단위를 갖추기 위한 장치이다.")
P("셋째, 이동 자료로 만든 권역은 계획을 평가할 때 쓸모 있는 비교 기준이지만, 그 자체가 최종 계획 단위는 아니다. 이 절의 결과는 이동 자료에 얼마나 잘 맞는지를 보여 줄 뿐, 공식 생활권을 대신하거나 이동 권역이 언제나 낫다는 근거가 아니다. "
  "Leiden과 Louvain의 비교도 이 자료와 조건에서의 결과이며, 두 방법의 일반적인 우열을 정하지 않는다.")
P(f"한계는 다음과 같다. (1) 1월 한 달, 서울 안의 통행이므로 계절, 요일, 서울 밖 목적지를 대표하지 않으며, HW·WH를 뺀 것은 제공 기관의 분류를 빌린 근사적인 기준이다. "
  f"(2) 3명 미만이라 공개되지 않은 값을 0으로 두었는데(조건에 맞는 행의 {p1(T1['masked_share_rows'])}%), 이것이 실제로 이동이 없었다는 뜻은 아니다. (3) 2023년 동 경계를 공통으로 쓰고 공식 생활권을 동 단위로 옮겼으므로, 동 안에서 경계가 갈리는 부분은 반영되지 않는다. "
  "(4) 구 경계와 권역 수를 고정하고 같은 자료로 권역을 만들고 평가했으므로, 이 권역 수가 가장 좋다거나 다른 도시에도 그대로 맞는다거나 경계를 바꾸면 효과가 있다는 근거는 되지 않는다. "
  "(5) 비교용 무작위 권역은 특정한 만드는 방식에 따른 결과이고, 구별 상관의 95% 범위도 구 사이의 공간적 의존을 고려한 것이 아니다. (6) 시설, 접근성, 서비스 용량은 다루지 않았으므로 생활권이 서비스를 충분히 갖추었는지는 결론으로 내지 않는다.")
P(f"정리하면, 공식 생활권과 이동 권역은 부분적으로 겹치고(IoU {n3(T2['iou_1to1'])}), IFR의 차이({pp2(T2['g'])}%p)는 판정이 바뀌는 통행의 비율({p2(T2['d_flow'])}%)보다 훨씬 작다. 같은 절차의 Leiden과 Louvain은 대부분의 구에서 같은 권역을 만들지만, 다시 계산했을 때 결과가 그대로 나오는 정도는 Leiden이 더 높았다. "
  "따라서 생활권 점검은 공간적 겹침, 권역 안에 담기는 이동의 차이와 방향, 판정이 바뀌는 통행의 규모, 결과가 다시 나오는 정도를 함께 보여 주고, 계획의 목적과 책임 체계 안에서 해석해야 한다. 이 절의 권역과 지표 정의는 4.3의 경계 조정 분석과 4.4의 두 시점 점검으로 이어진다.")

H(2, "부록 A. 구별 Louvain과 Leiden 결과")
rows = []
for r in L1.itertuples():
    rows.append([r.ku_name, f"{int(r.target)}", f"{r.louvain_resolution:.2f}", n4(r.louvain_q), p2(r.louvain_ifr), p1(r.louvain_modal_share), n3(r.louvain_stability_ari_min),
                 n4(r.leiden_q), p2(r.leiden_ifr), p1(r.leiden_modal_share), n3(r.leiden_louvain_ari)])
TBL("표 A-1", "구별 Louvain과 Leiden의 결과(2025년).", ["구", "목표 권역 수", "γ Louvain", "Q Louvain", "IFR Louvain (%)", "가장 잦은 결과 비율 Louvain (%)", "재계산 ARI 최솟값 Louvain", "Q Leiden", "IFR Leiden (%)", "가장 잦은 결과 비율 Leiden (%)", "두 결과의 ARI"], rows,
    "주: 재계산 ARI 최솟값은 고른 해상도에서 같은 계산을 10번 다시 했을 때 처음 결과와의 ARI 가운데 가장 작은 값이다. Leiden은 25개 구 모두 1이다.", small=True,
    widths=[1.4, 1.2, 1.3, 1.4, 1.5, 1.7, 1.6, 1.4, 1.5, 1.7, 1.3])
H(2, "부록 B. 자료와 결과의 연결")
P(f"이 절의 수치는 모두 결과표에서 만들었다. 자료 규모와 공식 생활권 구성은 T1과 C1, Louvain과 Leiden의 구별 비교는 L1과 L2, 공식 생활권과 Leiden의 비교, 지표 관계, 평가 방식을 바꾼 결과는 공통 분석 단계의 결과표 T2, T3, T6, T7, T8, A2에서 읽었다. "
  f"Louvain 권역은 {louv_info['year']}년 자료에서 시드 {louv_info['base_seed']}로 만들었고 계산에 {louv_info['seconds'] / 60:.1f}분이 걸렸다. 입력 파일의 SHA-256, 코드, 실행 환경은 패키지의 기록 파일에 있다. 본문의 값은 반올림했으며 CSV 결과표에는 더 자세한 값이 있다.")

H(2, "참고문헌")
for ref in [   # Elsevier Harvard 방식(JTG 게재본과 같은 표기). 웹 자료는 URL과 접속일
    "서울특별시, 2018. 2030 서울생활권계획. URL https://urban.seoul.go.kr/view/html/PMNU3040000001 (접속일: 2026.10.02).",
    "서울특별시, 2021. 서울 생활이동 데이터. 서울 열린데이터광장. URL http://data.seoul.go.kr/dataVisual/seoul/seoulLivingMigration.do (접속일: 2026.10.02).",
    "Blondel, V.D., Guillaume, J.-L., Lambiotte, R., Lefebvre, E., 2008. Fast unfolding of communities in large networks. J. Stat. Mech: Theory Exp. 2008 (10), P10008. https://doi.org/10.1088/1742-5468/2008/10/P10008.",
    "Hubert, L., Arabie, P., 1985. Comparing partitions. J. Classif. 2 (1), 193–218. https://doi.org/10.1007/BF01908075.",
    "Jaccard, P., 1912. The distribution of the flora in the alpine zone. New Phytol. 11 (2), 37–50. https://doi.org/10.1111/j.1469-8137.1912.tb05611.x.",
    "Lancichinetti, A., Fortunato, S., 2012. Consensus clustering in complex networks. Sci. Rep. 2, 336. https://doi.org/10.1038/srep00336.",
    "Newman, M.E.J., 2006. Modularity and community structure in networks. Proc. Natl. Acad. Sci. 103 (23), 8577–8582. https://doi.org/10.1073/pnas.0601602103.",
    "Park, J., Eom, S., Lee, M.-H., 2026. Benchmarking living-zone plans with mobility community detection: evidence from Seoul’s mobile-phone-based mobility data. J. Transp. Geogr. 135, 104753. https://doi.org/10.1016/j.jtrangeo.2026.104753.",
    "Reichardt, J., Bornholdt, S., 2006. Statistical mechanics of community detection. Phys. Rev. E 74 (1), 016110. https://doi.org/10.1103/PhysRevE.74.016110.",
    "Traag, V.A., Waltman, L., Van Eck, N.J., 2019. From Louvain to Leiden: guaranteeing well-connected communities. Sci. Rep. 9, 5233. https://doi.org/10.1038/s41598-019-41695-z.",
]:
    REF(ref)


# ═════════════════════════════════════ 출력: Markdown ═════════════════════════════════════
def md_out():
    out = []
    for b in blocks:
        if b[0] == "h":
            out.append("#" * b[1] + " " + b[2])
        elif b[0] in ("p", "r"):
            out.append(plain(b[1]))
        elif b[0] == "eq":
            out.append("`" + plain(b[1]) + "`  (" + str(b[2]) + ")")
        elif b[0] == "t":
            _, label, title, head, rows, note, *_ = b
            out.append(f"**{label}**  \n{plain(title)}")
            out.append("| " + " | ".join(head) + " |\n|" + "|".join(["---"] * len(head)) + "|\n" + "\n".join("| " + " | ".join(plain(str(c)) for c in r) + " |" for r in rows))
            if note:
                out.append(plain(note))
        elif b[0] == "f":
            _, name, label, title, _w = b
            out.append(f"![{label}](../results/figures/{name}.png)\n\n**{label}** {plain(title)}")
    (OUTD / "연구2_생활권_구획_점검_4.2.md").write_text("\n\n".join(out) + "\n", encoding="utf-8")


# ═════════════════════════════════════ 출력: DOCX ═════════════════════════════════════
FONT = "맑은 고딕"
MATH_LATIN = "Times New Roman"          # 수식의 영문·기호(한글은 맑은 고딕). 모든 수식이 같은 글꼴·같은 크기
EQ_PT = 10.5
TOKEN = re.compile(r"(_\{[^}]*\}|\^\{[^}]*\}|\*[^*\s][^*]*\*|D_flow|(?<![A-Za-z가-힣])[GQ](?![A-Za-z]))")


def minus(text):
    """음수 부호는 하이픈(-)이 아니라 마이너스 기호(−, U+2212)로 쓴다(그림 축 눈금과 같은 기호). 2025-01 같은 날짜·코드는 그대로 둔다."""
    return re.sub(r"(?<![0-9A-Za-z가-힣)\]])-(?=\d)", "−", text)


def plain(text):
    """Markdown용: 표기 규칙을 읽기 쉬운 평문으로 바꾼다."""
    t = re.sub(r"_\{([^}]*)\}", r"_\1", minus(text))
    t = re.sub(r"\^\{([^}]*)\}", r"^\1", t)
    return re.sub(r"\*([^*\s][^*]*)\*", r"\1", t)


def rich(par, text, size, bold=False, math=False):
    """표기 규칙: _{..} 아래 첨자, ^{..} 위 첨자, *..* 기울임. D_flow 와 홀로 쓰인 G·Q 는 기호(기울임)."""
    def add(seg, italic=False, sub=False, sup=False):
        r = par.add_run(seg)
        setfont(r, size, bold=bold, italic=italic)
        if math:
            r.font.name = MATH_LATIN
            r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        if sub:
            r.font.subscript = True
        if sup:
            r.font.superscript = True
    for part in TOKEN.split(minus(text)):
        if not part:
            continue
        if part == "D_flow":
            add("D", italic=True); add("flow", sub=True)
        elif part in ("G", "Q"):
            add(part, italic=True)
        elif part[:2] in ("_{", "^{"):
            for seg in re.split(r"(\*[^*]+\*)", part[2:-1]):
                if seg:
                    it = len(seg) > 2 and seg[0] == "*" and seg[-1] == "*"
                    add(seg[1:-1] if it else seg, italic=it, sub=part[0] == "_", sup=part[0] == "^")
        elif len(part) > 2 and part[0] == "*" and part[-1] == "*":
            add(part[1:-1], italic=True)
        else:
            add(part)
def setfont(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)


def _border(parent, tag, sz):
    el = OxmlElement(f"w:{tag}")
    if sz:
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
    else:
        el.set(qn("w:val"), "nil")
    parent.append(el)


def jtg_table_borders(tab):
    """표 전체: 위 선·아래 선만, 좌우·세로·안쪽 가로선 없음. 머리행 셀: 아래 선."""
    tblPr = tab._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for tag, sz in (("top", RULE_TOP), ("left", 0), ("bottom", RULE_BOTTOM), ("right", 0), ("insideH", 0), ("insideV", 0)):
        _border(b, tag, sz)
    tblPr.append(b)
    for cell in tab.rows[0].cells:
        tcPr = cell._element.get_or_add_tcPr()
        tcb = OxmlElement("w:tcBorders"); _border(tcb, "bottom", RULE_HEAD); tcPr.append(tcb)
    # 셀 안쪽 여백(위·아래 조금, 좌우 기본보다 작게)
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", 30), ("left", 70), ("bottom", 30), ("right", 70)):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), str(w)); el.set(qn("w:type"), "dxa"); mar.append(el)
    tblPr.append(mar)


def docx_out():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = Cm(2.4); sec.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]; normal.font.name = FONT; normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing = Pt(16); normal.paragraph_format.space_after = Pt(6)
    sp = normal.element.get_or_add_pPr().find(qn("w:spacing"))     # 한글·영문 사이 자동 간격 끔(스키마 순서상 spacing 앞)
    for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
        el = OxmlElement(tag); el.set(qn("w:val"), "0"); sp.addprevious(el)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        st = doc.styles[name]; st.font.name = FONT; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0x1F, 0x2D, 0x3D)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_before = Pt(14 if name != "Heading 3" else 10); st.paragraph_format.space_after = Pt(6); st.paragraph_format.keep_with_next = True
    text_w = 21.0 - 2 * 2.5
    for b in blocks:
        if b[0] == "h":
            doc.add_heading(b[2], level=b[1])
        elif b[0] == "p":
            para = doc.add_paragraph(); para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            rich(para, b[1], 10.5)
        elif b[0] == "eq":
            # 가운데 탭(본문 폭 절반)에 수식, 오른쪽 탭(본문 끝)에 번호. 줄 간격은 단일(첨자가 잘리지 않게)
            para = doc.add_paragraph(); pf = para.paragraph_format
            pf.space_before = Pt(4); pf.space_after = Pt(6); pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.tab_stops.add_tab_stop(Cm(text_w / 2), WD_TAB_ALIGNMENT.CENTER); pf.tab_stops.add_tab_stop(Cm(text_w), WD_TAB_ALIGNMENT.RIGHT)
            setfont(para.add_run("\t"), EQ_PT); rich(para, b[1], EQ_PT, math=True); setfont(para.add_run(f"\t({b[2]})"), EQ_PT)
        elif b[0] == "r":
            para = doc.add_paragraph(); para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.left_indent = Cm(0.6); para.paragraph_format.first_line_indent = Cm(-0.6)
            setfont(para.add_run(b[1]), 9.5)
        elif b[0] == "t":
            _, label, title, head, rows, note, small, widths, align = b
            # 캡션(표 위): 번호 줄(굵게) + 제목 줄(보통)
            c1 = doc.add_paragraph(); c1.paragraph_format.keep_with_next = True; c1.paragraph_format.space_before = Pt(10); c1.paragraph_format.space_after = Pt(0)
            setfont(c1.add_run(label), CAPTION_FONT_PT, bold=True)
            c2 = doc.add_paragraph(); c2.paragraph_format.keep_with_next = True; c2.paragraph_format.space_after = Pt(3)
            rich(c2, title, CAPTION_FONT_PT)
            tab = doc.add_table(rows=1, cols=len(head)); tab.alignment = WD_TABLE_ALIGNMENT.CENTER; tab.autofit = False
            fs = TABLE_FONT_SMALL_PT if small else TABLE_FONT_PT
            al = align or ["L"] * len(head)
            amap = {"L": WD_ALIGN_PARAGRAPH.LEFT, "C": WD_ALIGN_PARAGRAPH.CENTER, "R": WD_ALIGN_PARAGRAPH.RIGHT}
            for i, h in enumerate(head):
                cell = tab.rows[0].cells[i]; cell.text = ""
                p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05; p.alignment = amap[al[i]]
                rich(p, h, fs)
            for r in rows:
                cells = tab.add_row().cells
                for i, v in enumerate(r):
                    cells[i].text = ""; p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
                    p.alignment = amap[al[i]]
                    rich(p, str(v), fs)
            ws = widths or [text_w / len(head)] * len(head)
            for row in tab.rows:
                row.cells[0]._element.getparent()  # noqa (행 존재 확인)
                for i, w in enumerate(ws):
                    row.cells[i].width = Cm(w)
            trPr = tab.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)   # 쪽이 넘어가면 머리행 반복
            jtg_table_borders(tab)
            if note:
                n = doc.add_paragraph(); n.paragraph_format.space_before = Pt(3); n.paragraph_format.space_after = Pt(8)
                rich(n, note, NOTE_FONT_PT)
        elif b[0] == "f":
            _, name, label, title, width = b
            para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER; para.paragraph_format.keep_with_next = True; para.paragraph_format.space_before = Pt(8)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE      # 고정 줄 간격이면 그림이 한 줄 높이로 잘린다
            from PIL import Image                                                    # 그림을 그린 크기 그대로 넣어 글자 크기(7 pt 이상)를 지킨다
            with Image.open(FIG / f"{name}.png") as im:
                natural = im.size[0] / float(im.info.get("dpi", (300, 300))[0])
            para.add_run().add_picture(str(FIG / f"{name}.png"), width=Inches(min(natural, width)))
            c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(10)
            setfont(c.add_run(label + " "), CAPTION_FONT_PT, bold=True); rich(c, title, CAPTION_FONT_PT)
    doc.core_properties.title = "4.2 생활권 구획과 공식 생활권 점검"; doc.core_properties.author = ""
    doc.save(OUTD / "연구2_생활권_구획_점검_4.2.docx")


if __name__ == "__main__":
    md_out(); docx_out()
    n_tab = sum(1 for b in blocks if b[0] == "t"); n_fig = sum(1 for b in blocks if b[0] == "f")
    chars = sum(len(b[1]) for b in blocks if b[0] == "p")
    print(f"원고 생성: 표 {n_tab}, 그림 {n_fig}, 본문 {chars:,}자 → {OUTD}")
