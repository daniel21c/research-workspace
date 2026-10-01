# -*- coding: utf-8 -*-
"""별도 해제 폴더에서 분석을 처음부터 다시 실행하고, 저장된 결과와 대조한다. 공통 코어엔진·허브 폴더를 읽지 않는다.

  python code/reproduce.py            # S1·S4·S5(그림·원고)를 다시 만들어 저장된 결과와 대조 (수 분)
  python code/reproduce.py --full     # 위에 더해 S3(Louvain 2025, 25구 × 250γ × 3,000회)를 처음부터 다시 실행해 대조 (워커 25개로 약 26분)

대조 규칙: 표(CSV)·원고 MD는 바이트 일치, 그림(PNG)은 바이트 일치(불일치하면 픽셀 일치), DOCX는 문단·표 글자 일치(압축 시각이 달라 바이트 비교 불가).
PDF는 Word로만 출력하므로(code/export_pdf.ps1) 이 스크립트에서 다시 만들지 않는다.
결과: audit/reproduction_report.json
"""
from __future__ import annotations
import argparse, hashlib, json, os, platform, shutil, subprocess, sys, time
from pathlib import Path
import numpy as np

PKG = Path(__file__).resolve().parents[1]
PY = sys.executable
REF = PKG / "audit" / "_reference_copy"
TRACKED = ["results/tables", "results/figures", "manuscript"]
report = {"started": time.strftime("%Y-%m-%dT%H:%M:%S"), "python": sys.version.split()[0], "platform": platform.platform(), "steps": {}}


def sha(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""): h.update(b)
    return h.hexdigest()


def run(label, args, env=None):
    t0 = time.time()
    r = subprocess.run([PY, *args], cwd=PKG, capture_output=True, text=True, encoding="utf-8", env={**os.environ, "PYTHONIOENCODING": "utf-8", **(env or {})})
    report["steps"][label] = {"returncode": r.returncode, "seconds": round(time.time() - t0, 1), "stdout_tail": r.stdout[-400:], "stderr_tail": r.stderr[-400:] if r.returncode else ""}
    print(f"[{label}] rc={r.returncode} {time.time() - t0:.0f}s", flush=True)
    if r.returncode:
        print(r.stderr[-1500:]); finish(False)
    return r


def snapshot():
    if REF.exists():
        shutil.rmtree(REF)
    for rel in TRACKED + ["audit/s1_checks.json", "audit/s4_checks.json"]:
        src = PKG / rel
        dst = REF / rel
        if src.is_dir():
            shutil.copytree(src, dst, ignore=shutil.ignore_patterns("*.pdf"))
        elif src.exists():
            dst.parent.mkdir(parents=True, exist_ok=True); shutil.copyfile(src, dst)


def docx_text(p):
    from docx import Document
    d = Document(p)
    return [x.text for x in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]


def compare():
    from PIL import Image
    res = {"identical": [], "pixel_identical_only": [], "docx_text_identical": [], "different": []}
    for f in sorted(REF.rglob("*")):
        if not f.is_file():
            continue
        rel = f.relative_to(REF); now = PKG / rel
        if not now.exists():
            res["different"].append(f"{rel} (missing)"); continue
        if f.suffix == ".docx":
            (res["docx_text_identical"] if docx_text(f) == docx_text(now) else res["different"]).append(str(rel))
        elif sha(f) == sha(now):
            res["identical"].append(str(rel))
        elif f.suffix == ".png" and np.array_equal(np.asarray(Image.open(f)), np.asarray(Image.open(now))):
            res["pixel_identical_only"].append(str(rel))
        else:
            res["different"].append(str(rel))
    return res


def compare_s3(rerun: Path):
    """S3 재실행(output/louvain/2025_repro)과 저장된 결과(output/louvain/2025)를 대조: 매핑 CSV 바이트, 확정 해상도 3,000회 원시 라벨 전부."""
    orig = PKG / "output" / "louvain" / "2025"
    out = {"mapping_csv_identical": sha(orig / "metrics" / "louvain_mapping_2025.csv") == sha(rerun / "metrics" / "louvain_mapping_2025.csv")}
    n_ok = n = 0
    for f in sorted((orig / "coassoc").glob("*_runs.npz")):
        a, b = np.load(f), np.load(rerun / "coassoc" / f.name)
        n += 1; n_ok += int(np.array_equal(a["runs"], b["runs"]) and np.array_equal(a["labels_cc"], b["labels_cc"]))
    out["raw_label_files_identical"] = f"{n_ok}/{n}"
    sc_ok = sum(1 for f in (orig / "resolution_scan_logs").glob("*.csv") if sha(f) == sha(rerun / "resolution_scan_logs" / f.name) or
                (lambda x, y: x.drop(columns=[c for c in x.columns if c == "seconds"]).equals(y.drop(columns=[c for c in y.columns if c == "seconds"])))(
                    __import__("pandas").read_csv(f), __import__("pandas").read_csv(rerun / "resolution_scan_logs" / f.name)))
    out["scan_logs_identical"] = f"{sc_ok}/{len(list((orig / 'resolution_scan_logs').glob('*.csv')))}"
    out["pass"] = bool(out["mapping_csv_identical"] and n_ok == n and sc_ok == len(list((orig / "resolution_scan_logs").glob("*.csv"))))
    return out


def finish(ok):
    report["finished"] = time.strftime("%Y-%m-%dT%H:%M:%S"); report["all_pass"] = bool(ok)
    (PKG / "audit" / "reproduction_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print("ALL PASS" if ok else "FAILED"); sys.exit(0 if ok else 1)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--full", action="store_true"); ap.add_argument("--workers", type=int, default=25)
    a = ap.parse_args()
    run("verify_inputs", ["code/verify_inputs.py"])
    snapshot()
    louv_env = None
    if a.full:
        rerun = PKG / "output" / "louvain" / "2025_repro"
        if rerun.exists():
            shutil.rmtree(rerun)
        run("s3_louvain_rerun", ["code/engine/s03_louvain_consensus.py", "--years", "2025", "--workers", str(a.workers), "--tag", "_repro"])
        report["s3_comparison"] = compare_s3(rerun)
        if not report["s3_comparison"]["pass"]:
            finish(False)
        louv_env = {"LOUVAIN_DIR": str(rerun)}
    run("s1_overview", ["code/s1_overview.py"])
    run("s4_compare", ["code/s4_compare.py"], louv_env)
    run("s5_figures", ["code/s5_figures.py"])
    run("s5_build_manuscript", ["code/s5_build_manuscript.py"])
    report["comparison"] = compare()
    report["comparison"]["n_compared"] = sum(len(v) for v in report["comparison"].values())
    ok = not report["comparison"]["different"]
    shutil.rmtree(REF, ignore_errors=True)
    finish(ok)


if __name__ == "__main__":
    main()
